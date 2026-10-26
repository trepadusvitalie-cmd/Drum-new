"""Coordonate reale și GIS: proiecția MOLDREF99 / Moldova TM (EPSG:4026) ↔ WGS84,
export SHP / GDB / XLS / XLSX / KML / KMZ, import KML / KMZ / GPX („Задать местоположение → Из файла”)
și harta de fundal Esri („Из карт Esri”), descărcată în fundal și ținută în cache pe disc."""
import io, math, os, re, struct, threading, zipfile
import tkinter as tk
from tkinter import filedialog, messagebox
from xml.sax.saxutils import escape
from stil import *

# ---------------------------------------------------------------- proiecția (Transverse Mercator pe GRS80)
A_ELIPS, F_ELIPS = 6378137.0, 1 / 298.257222101
SISTEME = {"MOLDREF99 / Moldova TM (EPSG:4026)": dict(lon0=28.4, k0=0.99994, fe=200000.0, fn=-5000000.0, epsg=4026)}
SISTEM = "MOLDREF99 / Moldova TM (EPSG:4026)"
PRJ_4026 = ('PROJCS["MOLDREF99_Moldova_TM",GEOGCS["GCS_MOLDREF99",DATUM["D_MOLDREF99",SPHEROID["GRS_1980",6378137.0,298.257222101]],'
            'PRIMEM["Greenwich",0.0],UNIT["Degree",0.0174532925199433]],PROJECTION["Transverse_Mercator"],PARAMETER["False_Easting",200000.0],'
            'PARAMETER["False_Northing",-5000000.0],PARAMETER["Central_Meridian",28.4],PARAMETER["Scale_Factor",0.99994],'
            'PARAMETER["Latitude_Of_Origin",0.0],UNIT["Meter",1.0]]')

_e2 = F_ELIPS * (2 - F_ELIPS); _ep2 = _e2 / (1 - _e2)
_m1 = 1 - _e2 / 4 - 3 * _e2 ** 2 / 64 - 5 * _e2 ** 3 / 256
_m2 = 3 * _e2 / 8 + 3 * _e2 ** 2 / 32 + 45 * _e2 ** 3 / 1024
_m3 = 15 * _e2 ** 2 / 256 + 45 * _e2 ** 3 / 1024
_m4 = 35 * _e2 ** 3 / 3072


def _arc_meridian(f):
    return A_ELIPS * (_m1 * f - _m2 * math.sin(2 * f) + _m3 * math.sin(4 * f) - _m4 * math.sin(6 * f))


def din_wgs(lat, lon, s=SISTEM):
    """grade WGS84 → (X, Y) în metri MOLDREF99"""
    p = SISTEME[s]; f = math.radians(lat); l = math.radians(lon - p["lon0"])
    sn, cs = math.sin(f), math.cos(f); n = A_ELIPS / math.sqrt(1 - _e2 * sn * sn); t = math.tan(f) ** 2; c = _ep2 * cs * cs; a = l * cs
    x = p["k0"] * n * (a + (1 - t + c) * a ** 3 / 6 + (5 - 18 * t + t * t + 72 * c - 58 * _ep2) * a ** 5 / 120)
    y = p["k0"] * (_arc_meridian(f) + n * math.tan(f) * (a * a / 2 + (5 - t + 9 * c + 4 * c * c) * a ** 4 / 24
                                                       + (61 - 58 * t + t * t + 600 * c - 330 * _ep2) * a ** 6 / 720))
    return p["fe"] + x, p["fn"] + y


def in_wgs(x, y, s=SISTEM):
    """(X, Y) MOLDREF99 → (lat, lon) grade WGS84"""
    p = SISTEME[s]; m = (y - p["fn"]) / p["k0"]; mu = m / (A_ELIPS * _m1)
    e1 = (1 - math.sqrt(1 - _e2)) / (1 + math.sqrt(1 - _e2))
    f1 = (mu + (3 * e1 / 2 - 27 * e1 ** 3 / 32) * math.sin(2 * mu) + (21 * e1 ** 2 / 16 - 55 * e1 ** 4 / 32) * math.sin(4 * mu)
          + (151 * e1 ** 3 / 96) * math.sin(6 * mu) + (1097 * e1 ** 4 / 512) * math.sin(8 * mu))
    sn, cs = math.sin(f1), math.cos(f1); c1 = _ep2 * cs * cs; t1 = math.tan(f1) ** 2
    n1 = A_ELIPS / math.sqrt(1 - _e2 * sn * sn); r1 = A_ELIPS * (1 - _e2) / (1 - _e2 * sn * sn) ** 1.5; d = (x - p["fe"]) / (n1 * p["k0"])
    lat = f1 - (n1 * math.tan(f1) / r1) * (d * d / 2 - (5 + 3 * t1 + 10 * c1 - 4 * c1 * c1 - 9 * _ep2) * d ** 4 / 24
                                           + (61 + 90 * t1 + 298 * c1 + 45 * t1 * t1 - 252 * _ep2 - 3 * c1 * c1) * d ** 6 / 720)
    lon = (d - (1 + 2 * t1 + c1) * d ** 3 / 6 + (5 - 2 * c1 + 28 * t1 - 3 * c1 * c1 + 8 * _ep2 + 24 * t1 * t1) * d ** 5 / 120) / cs
    return math.degrees(lat), p["lon0"] + math.degrees(lon)


# ---------------------------------------------------------------- obiectele desenului ca geometrii GIS
def colecteaza(doc, handles=None):
    """listă de dict(tip=Point/LineString/Polygon, xy=[(x,y)…], strat, obiect, handle, bloc, text, atribute, lungime, arie)"""
    from ezdxf import path as ezpath
    rez = []
    for e in doc.modelspace():
        h = e.dxf.handle
        if handles is not None and h not in handles: continue
        t = e.dxftype(); r = dict(strat=e.dxf.get("layer", "0"), obiect=t, handle=h, bloc="", text="", atribute={})
        try:
            if t == "INSERT":
                p = e.dxf.insert; r.update(tip="Point", xy=[(p.x, p.y)], bloc=e.dxf.name, atribute={a.dxf.tag: a.dxf.text for a in e.attribs})
            elif t in ("TEXT", "MTEXT"):
                p = e.dxf.insert; r.update(tip="Point", xy=[(p.x, p.y)], text=e.plain_text() if t == "MTEXT" else e.dxf.text)
            elif t == "POINT":
                p = e.dxf.location; r.update(tip="Point", xy=[(p.x, p.y)])
            elif t == "HATCH":
                for pa in ezpath.from_hatch(e):
                    pts = [(v.x, v.y) for v in pa.flattening(0.05)]
                    if len(pts) >= 3: rez.append(dict(r, tip="Polygon", xy=pts))
                continue
            elif t in ("LINE", "LWPOLYLINE", "POLYLINE", "ARC", "CIRCLE", "ELLIPSE", "SPLINE"):
                pts = [(v.x, v.y) for v in ezpath.make_path(e).flattening(0.05)]
                if len(pts) < 2: continue
                inchis = t == "CIRCLE" or (t in ("LWPOLYLINE", "POLYLINE") and e.is_closed) or (len(pts) > 3 and pts[0] == pts[-1])
                r.update(tip="Polygon" if inchis and len(pts) >= 3 else "LineString", xy=pts)
            else: continue
        except Exception: continue
        rez.append(r)
    for r in rez:
        xy = r["xy"]
        r["lungime"] = sum(math.hypot(xy[i + 1][0] - xy[i][0], xy[i + 1][1] - xy[i][1]) for i in range(len(xy) - 1)) if r["tip"] != "Point" else 0.0
        r["arie"] = abs(sum(xy[i][0] * xy[(i + 1) % len(xy)][1] - xy[(i + 1) % len(xy)][0] * xy[i][1] for i in range(len(xy)))) / 2 if r["tip"] == "Polygon" else 0.0
    return rez


def _campuri_atribute(ob):
    tags = []
    for r in ob:
        for k in r["atribute"]:
            if k not in tags: tags.append(k)
    return tags


def _wkb(r):
    xy = r["xy"]
    if r["tip"] == "Point": return struct.pack("<BIdd", 1, 1, *xy[0])
    if r["tip"] == "LineString": return struct.pack("<BII", 1, 2, len(xy)) + b"".join(struct.pack("<dd", *p) for p in xy)
    if xy[0] != xy[-1]: xy = xy + [xy[0]]
    return struct.pack("<BIII", 1, 3, 1, len(xy)) + b"".join(struct.pack("<dd", *p) for p in xy)


NUME_STRAT = {"Point": "puncte", "LineString": "linii", "Polygon": "poligoane"}


def export_shp(ob, cale):
    import shapefile
    baza = os.path.splitext(cale)[0]; tags = _campuri_atribute(ob); scrise = []
    scurte = {}; folosite = {"LAYER", "OBJECT", "HANDLE", "BLOCK", "TEXT", "LENGTH", "AREA"}
    for t in tags:                                          # numele câmpurilor în DBF: max. 10 caractere, unice
        n = re.sub(r"[^A-Za-z0-9_]", "_", t)[:10].upper() or "ATTR"; k = 1
        while n in folosite: n = n[:8] + str(k); k += 1
        folosite.add(n); scurte[t] = n
    for tip, tipshp in (("Point", shapefile.POINT), ("LineString", shapefile.POLYLINE), ("Polygon", shapefile.POLYGON)):
        rows = [r for r in ob if r["tip"] == tip]
        if not rows: continue
        f = f"{baza}_{NUME_STRAT[tip]}"
        with shapefile.Writer(f, shapeType=tipshp, encoding="utf-8") as w:
            for n, tp, ln in (("LAYER", "C", 80), ("OBJECT", "C", 20), ("HANDLE", "C", 16), ("BLOCK", "C", 80), ("TEXT", "C", 254)):
                w.field(n, tp, ln)
            w.field("LENGTH", "N", 18, 3); w.field("AREA", "N", 18, 3)
            for t in tags: w.field(scurte[t], "C", 254)
            for r in rows:
                if tip == "Point": w.point(*r["xy"][0])
                elif tip == "LineString": w.line([r["xy"]])
                else:
                    ring = r["xy"] if r["xy"][0] == r["xy"][-1] else r["xy"] + [r["xy"][0]]
                    if _arie_semn(ring) > 0: ring = ring[::-1]          # shapefile: inel exterior în sensul acelor de ceasornic
                    w.poly([ring])
                w.record(r["strat"], r["obiect"], r["handle"], r["bloc"], r["text"][:254], round(r["lungime"], 3), round(r["arie"], 3),
                         *[str(r["atribute"].get(t, ""))[:254] for t in tags])
        open(f + ".prj", "w").write(PRJ_4026); open(f + ".cpg", "w").write("UTF-8")
        scrise.append(os.path.basename(f) + ".shp")
    return scrise


def _arie_semn(xy):
    return sum(xy[i][0] * xy[i + 1][1] - xy[i + 1][0] * xy[i][1] for i in range(len(xy) - 1)) / 2


def export_gdb(ob, cale):
    import numpy as np
    from pyogrio.raw import write
    if not cale.lower().endswith(".gdb"): cale += ".gdb"
    tags = _campuri_atribute(ob); scrise = []
    for tip in ("Point", "LineString", "Polygon"):
        rows = [r for r in ob if r["tip"] == tip]
        if not rows: continue
        campuri = ["LAYER", "OBJECT", "HANDLE", "BLOCK", "TEXT", "LENGTH", "AREA"] + [re.sub(r"[^A-Za-z0-9_]", "_", t) or "ATTR" for t in tags]
        date = [np.array([r["strat"] for r in rows], dtype=object), np.array([r["obiect"] for r in rows], dtype=object),
                np.array([r["handle"] for r in rows], dtype=object), np.array([r["bloc"] for r in rows], dtype=object),
                np.array([r["text"] for r in rows], dtype=object), np.array([r["lungime"] for r in rows]), np.array([r["arie"] for r in rows])]
        date += [np.array([str(r["atribute"].get(t, "")) for r in rows], dtype=object) for t in tags]
        geom = np.array([_wkb(r) for r in rows], dtype=object)
        write(cale, geometry=geom, field_data=date, fields=campuri, layer=NUME_STRAT[tip], driver="OpenFileGDB",
              geometry_type=tip, crs="EPSG:4026", encoding="UTF-8", append=False)
        scrise.append(os.path.basename(cale) + " / " + NUME_STRAT[tip])
    return scrise


def _randuri_tabel(ob):
    tags = _campuri_atribute(ob)
    cap = ["HANDLE", "Объект", "Слой", "X", "Y", "Блок", "Текст", "Длина", "Площадь", "Вершин"] + tags
    rows = [[r["handle"], r["obiect"], r["strat"], round(r["xy"][0][0], 4), round(r["xy"][0][1], 4), r["bloc"], r["text"],
             round(r["lungime"], 3), round(r["arie"], 3), len(r["xy"])] + [r["atribute"].get(t, "") for t in tags] for r in ob]
    return cap, rows


def export_xlsx(ob, cale):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    cap, rows = _randuri_tabel(ob); wb = Workbook(); ws = wb.active; ws.title = "Объекты"; ws.append(cap)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="3A5A8C")
    for r in rows: ws.append(r)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate([10, 14, 22, 14, 14, 18, 30, 12, 12, 8], 1): ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    wb.save(cale); return [os.path.basename(cale)]


def export_xls(ob, cale):
    import xlwt
    cap, rows = _randuri_tabel(ob)
    if len(rows) > 65534: raise ValueError(f"Формат XLS вмещает максимум 65534 строки, а объектов {len(rows)}. Сохраните как XLSX.")
    wb = xlwt.Workbook(encoding="utf-8"); ws = wb.add_sheet("Объекты"); st = xlwt.easyxf("font: bold on")
    for j, v in enumerate(cap[:256]): ws.write(0, j, v, st)
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r[:256]): ws.write(i, j, v)
    wb.save(cale); return [os.path.basename(cale)]


def _kml(ob, nume):
    from ezdxf import colors as ezcolors
    straturi = {}
    for r in ob: straturi.setdefault(r["strat"], []).append(r)
    out = ['<?xml version="1.0" encoding="UTF-8"?>', '<kml xmlns="http://www.opengis.net/kml/2.2"><Document>', f"<name>{escape(nume)}</name>"]
    for s, rows in straturi.items():
        out.append(f"<Folder><name>{escape(s)}</name>")
        for r in rows:
            ll = [in_wgs(x, y) for x, y in r["xy"]]; c = " ".join(f"{lo:.8f},{la:.8f},0" for la, lo in ll)
            titlu = r["text"] or r["bloc"] or r["obiect"]
            descr = "".join(f"<b>{escape(k)}</b>: {escape(str(v))}<br/>" for k, v in r["atribute"].items())
            out.append(f"<Placemark><name>{escape(titlu)}</name><description><![CDATA[Слой: {escape(s)}<br/>HANDLE: {r['handle']}<br/>{descr}]]></description>")
            if r["tip"] == "Point": out.append(f"<Point><coordinates>{c}</coordinates></Point>")
            elif r["tip"] == "LineString": out.append(f"<Style><LineStyle><color>ff00ffff</color><width>2</width></LineStyle></Style><LineString><tessellate>1</tessellate><coordinates>{c}</coordinates></LineString>")
            else:
                if ll[0] != ll[-1]: c += f" {ll[0][1]:.8f},{ll[0][0]:.8f},0"
                out.append(f"<Style><PolyStyle><color>4000ff00</color></PolyStyle></Style><Polygon><outerBoundaryIs><LinearRing><coordinates>{c}</coordinates></LinearRing></outerBoundaryIs></Polygon>")
            out.append("</Placemark>")
        out.append("</Folder>")
    out.append("</Document></kml>")
    return "\n".join(out)


def export_kml(ob, cale):
    open(cale, "w", encoding="utf-8").write(_kml(ob, os.path.basename(cale))); return [os.path.basename(cale)]


def export_kmz(ob, cale):
    with zipfile.ZipFile(cale, "w", zipfile.ZIP_DEFLATED) as z: z.writestr("doc.kml", _kml(ob, os.path.basename(cale)))
    return [os.path.basename(cale)]


FORMATE_EXPORT = [("Shapefile (*.shp)", ".shp", export_shp), ("Esri File Geodatabase (*.gdb)", ".gdb", export_gdb),
                  ("Excel (*.xlsx)", ".xlsx", export_xlsx), ("Excel 97-2003 (*.xls)", ".xls", export_xls),
                  ("Google Earth KML (*.kml)", ".kml", export_kml), ("Google Earth KMZ (*.kmz)", ".kmz", export_kmz)]


def dialog_export(app, fmt=None):
    """ЭКСПОРТГИС: alegi formatul, apoi fișierul; din selecție (dacă există) sau din tot desenul"""
    if not app.doc: return app.mesaj("Нет открытого чертежа.")
    if fmt is None:
        m = tk.Menu(app.root, tearoff=0, bg="#3b4453", fg=TEXT, activebackground=HOVER, font=(F, -13))
        for t, ext, _ in FORMATE_EXPORT: m.add_command(label="  " + t, command=lambda e=ext: dialog_export(app, e))
        x, y = app.root.winfo_pointerx(), app.root.winfo_pointery(); return m.tk_popup(x, y)
    nume, ext, f = next(x for x in FORMATE_EXPORT if x[1] == fmt)
    baza = os.path.splitext(os.path.basename(app.cale or "Чертеж1"))[0]
    if ext == ".gdb":
        d = filedialog.askdirectory(title="Папка, где создать базу .gdb")
        if not d: return
        cale = os.path.join(d, baza + ".gdb")
    else:
        cale = filedialog.asksaveasfilename(title="Экспорт: " + nume, initialfile=baza + ext, defaultextension=ext, filetypes=[(nume, "*" + ext)])
        if not cale: return
    ob = colecteaza(app.doc, set(app.ed.sel) if app.ed.sel else None)
    if not ob: return app.mesaj("Нет объектов для экспорта (точки, линии, полилинии, круги, блоки, тексты).")
    app.root.config(cursor="watch"); app.root.update_idletasks()
    try: scrise = f(ob, cale)
    except Exception as ex:
        import traceback; traceback.print_exc(); return messagebox.showerror("Экспорт", f"Не удалось экспортировать:\n{ex}")
    finally: app.root.config(cursor="")
    app.mesaj(f"Экспорт {nume}: {len(ob)} объектов ({'выбранные' if app.ed.sel else 'весь чертеж'}), система {SISTEM} → " + ", ".join(scrise))


# ---------------------------------------------------------------- import KML / KMZ / GPX în coordonate MOLDREF99
def _coordonate(t):
    rez = []
    for g in t.split():
        p = g.split(",")
        if len(p) >= 2:
            try: rez.append((float(p[1]), float(p[0])))
            except ValueError: pass
    return rez


def import_geo(app, cale=None):
    import xml.etree.ElementTree as ET
    if not app.doc: return app.mesaj("Сначала откройте или создайте чертеж.")
    cale = cale or filedialog.askopenfilename(title="Задать местоположение – из файла",
                                              filetypes=[("KML / KMZ / GPX", "*.kml *.kmz *.gpx"), ("Все файлы", "*.*")])
    if not cale: return
    try:
        if cale.lower().endswith(".kmz"):
            with zipfile.ZipFile(cale) as z: txt = z.read(next(n for n in z.namelist() if n.lower().endswith(".kml")))
        else: txt = open(cale, "rb").read()
        rad = ET.fromstring(txt)
    except Exception as ex: return messagebox.showerror("Местоположение", f"Не удалось прочитать файл:\n{ex}")
    loc = lambda el: el.tag.split("}")[-1]
    msp = app.doc.modelspace(); app.ed.instantaneu([]); n = 0; xs = []
    def strat(nume):
        nume = re.sub(r'[<>/\\":;?*|=`]', "_", nume or "GEO")[:200] or "GEO"
        if nume not in app.doc.layers: app.doc.layers.add(nume, color=3)
        return nume
    def adauga(tipul, pts, nume, s):
        nonlocal n
        xy = [din_wgs(la, lo) for la, lo in pts]; xs.extend(xy)
        if tipul == "p":
            msp.add_point(xy[0], dxfattribs={"layer": s})
            if nume: msp.add_text(nume, height=2.0, dxfattribs={"layer": s}).set_placement((xy[0][0] + 1, xy[0][1] + 1))
        elif len(xy) >= 2: msp.add_lwpolyline(xy, close=(tipul == "g"), dxfattribs={"layer": s})
        n += 1
    if loc(rad) == "gpx":
        for el in rad.iter():
            t = loc(el)
            if t == "wpt": adauga("p", [(float(el.get("lat")), float(el.get("lon")))], next((c.text for c in el if loc(c) == "name"), ""), strat("GPX_точки"))
            elif t in ("trkseg", "rte"):
                pts = [(float(c.get("lat")), float(c.get("lon"))) for c in el if loc(c) in ("trkpt", "rtept")]
                if pts: adauga("l", pts, "", strat("GPX_треки"))
    else:
        def parcurge(el, folder):
            for c in el:
                t = loc(c)
                if t in ("Folder", "Document"):
                    nm = next((x.text for x in c if loc(x) == "name"), None); parcurge(c, nm or folder)
                elif t == "Placemark":
                    nm = next((x.text for x in c if loc(x) == "name"), "") or ""; s = strat(folder or "KML")
                    for g in c.iter():
                        tg = loc(g)
                        if tg in ("Point", "LineString", "LinearRing"):
                            co = next((x.text for x in g.iter() if loc(x) == "coordinates"), "")
                            adauga({"Point": "p", "LineString": "l", "LinearRing": "g"}[tg], _coordonate(co or ""), nm, s)
                else: parcurge(c, folder)
        parcurge(rad, None)
    if not n: return app.mesaj("В файле нет точек или линий.")
    app.ed.gata_modificare(); app.marcheaza()
    x0, y0 = min(p[0] for p in xs), min(p[1] for p in xs); x1, y1 = max(p[0] for p in xs), max(p[1] for p in xs)
    pad = max(x1 - x0, y1 - y0, 50) * 0.1; app.v.zoom_fereastra(x0 - pad, y0 - pad, x1 + pad, y1 + pad)
    app.mesaj(f"Местоположение из файла {os.path.basename(cale)}: {n} объектов в системе {SISTEM}.")


# ---------------------------------------------------------------- harta de fundal (Esri)
SURSE = {
    "Спутник (Esri World Imagery)": "https://server.arcgisonline.com/ArcGIS/rest/services/World_Imagery/MapServer/tile/{z}/{y}/{x}",
    "Улицы (Esri World Street Map)": "https://server.arcgisonline.com/ArcGIS/rest/services/World_Street_Map/MapServer/tile/{z}/{y}/{x}",
    "Топография (Esri World Topo Map)": "https://server.arcgisonline.com/ArcGIS/rest/services/World_Topo_Map/MapServer/tile/{z}/{y}/{x}",
}
ATRIBUIRE = "Карта: Esri, Maxar, Earthstar Geographics, участники OpenStreetMap"


def _dosar_cache():
    d = os.path.join(os.environ.get("APPDATA") or os.path.expanduser("~"), "CAD_PC", "harti"); os.makedirs(d, exist_ok=True); return d


def _tile_xy(lat, lon, z):
    n = 2 ** z; x = (lon + 180) / 360 * n
    y = (1 - math.log(math.tan(math.radians(lat)) + 1 / math.cos(math.radians(lat))) / math.pi) / 2 * n
    return x, y


def _tile_ll(x, y, z):
    n = 2 ** z; lon = x / n * 360 - 180; lat = math.degrees(math.atan(math.sinh(math.pi * (1 - 2 * y / n))))
    return lat, lon


class Harta:
    """plăcile hărții sub desen: pentru fiecare placă vizibilă calculez colțurile în MOLDREF99 și o întind pe ecran"""
    def __init__(self, v):
        self.v = v; self.sursa = None; self.c = {}; self.in_lucru = set(); self.foto = []; self._verif = None; self.lock = threading.Lock()

    def porneste(self, sursa): self.sursa = sursa; self.v.deseneaza()
    def opreste(self): self.sursa = None; self.v.cv.delete("harta"); self.v.deseneaza()

    def _ia(self, z, x, y):
        k = (self.sursa, z, x, y)
        if k in self.c: return self.c[k]
        if k in self.in_lucru: return None
        self.in_lucru.add(k)
        def lucru():
            from PIL import Image
            import urllib.request
            f = os.path.join(_dosar_cache(), f"{list(SURSE).index(self.sursa)}_{z}_{x}_{y}.jpg"); im = None
            try:
                if os.path.exists(f): im = Image.open(f); im.load()
                else:
                    req = urllib.request.Request(SURSE[self.sursa].format(z=z, x=x, y=y), headers={"User-Agent": "CAD-PC/3.3"})
                    data = urllib.request.urlopen(req, timeout=15).read(); open(f, "wb").write(data); im = Image.open(io.BytesIO(data)); im.load()
                im = im.convert("RGB")
            except Exception: im = False
            with self.lock:
                if len(self.c) > 400: self.c.clear()
                self.c[k] = im; self.in_lucru.discard(k)
        threading.Thread(target=lucru, daemon=True).start()
        return None

    def deseneaza(self, W, H):
        cv = self.v.cv; cv.delete("harta"); self.foto = []
        if not self.sursa: return
        from PIL import Image, ImageTk
        v = self.v
        try:
            colt = [in_wgs(*v.lume(sx, sy)) for sx, sy in ((0, 0), (W, 0), (0, H), (W, H))]
        except Exception: return
        lats = [c[0] for c in colt]; lons = [c[1] for c in colt]
        if not (-85 < min(lats) and max(lats) < 85): return
        lat_c = sum(lats) / 4
        z = int(round(math.log2(max(1e-9, 156543.03 * math.cos(math.radians(lat_c)) * v.s))))
        z = max(2, min(19, z))
        while True:
            x0, y0 = _tile_xy(max(lats), min(lons), z); x1, y1 = _tile_xy(min(lats), max(lons), z)
            if (int(x1) - int(x0) + 1) * (int(y1) - int(y0) + 1) <= 64 or z <= 2: break
            z -= 1
        lipsa = False
        for ty in range(int(y0), int(y1) + 1):
            for tx in range(int(x0), int(x1) + 1):
                im = self._ia(z, tx, ty)
                if im is None: lipsa = True; continue
                if im is False: continue
                (la0, lo0), (la1, lo1) = _tile_ll(tx, ty, z), _tile_ll(tx + 1, ty + 1, z)
                sx0, sy0 = v.ecran(*din_wgs(la0, lo0)); sx1, sy1 = v.ecran(*din_wgs(la1, lo1))
                w, h = int(round(sx1 - sx0)), int(round(sy1 - sy0))
                if w < 1 or h < 1 or w > 4000 or h > 4000: continue
                f = ImageTk.PhotoImage(im.resize((w, h), Image.BILINEAR)); self.foto.append(f)
                cv.create_image(sx0, sy0, image=f, anchor="nw", tags=("harta",))
        cv.create_text(W - 8, H - 4, text=ATRIBUIRE, anchor="se", fill="#e8ebef", font=(F, -10), tags=("harta",))
        cv.tag_lower("harta")
        if lipsa and not self._verif: self._verif = cv.after(300, self._reverifica)

    def _reverifica(self):
        self._verif = None
        if self.sursa: self.v.deseneaza()


def meniu_esri(app):
    if not app.doc: return app.mesaj("Нет открытого чертежа.")
    d = tk.Toplevel(app.root); d.title("Задать местоположение – карты Esri"); d.configure(bg=CORP); d.transient(app.root)
    tk.Label(d, text="Система координат чертежа:", bg=CORP, fg=TEXT, font=F_MIC).grid(row=0, column=0, sticky="w", padx=10, pady=(10, 2))
    tk.Label(d, text=SISTEM, bg=CORP, fg="#8fd18f", font=F_MIC).grid(row=1, column=0, sticky="w", padx=20)
    tk.Label(d, text="Карта под чертежом:", bg=CORP, fg=TEXT, font=F_MIC).grid(row=2, column=0, sticky="w", padx=10, pady=(10, 2))
    v = tk.StringVar(value=app.v.harta.sursa or list(SURSE)[0])
    for i, n in enumerate(list(SURSE) + ["Выкл (без карты)"]):
        tk.Radiobutton(d, text=n, value=n, variable=v, bg=CORP, fg=TEXT, selectcolor="#2b313a", activebackground=CORP, font=F_MIC).grid(row=3 + i, column=0, sticky="w", padx=20)
    tk.Label(d, text="Нужен интернет. Карта только показывается, в DWG не сохраняется.\n" + ATRIBUIRE, bg=CORP, fg=GRI, font=(F, -11), justify="left").grid(row=8, column=0, sticky="w", padx=10, pady=6)
    def ok(e=None):
        n = v.get(); d.destroy()
        if n in SURSE: app.v.harta.porneste(n); app.mesaj(f"Карта: {n}. Колесо мыши – масштаб, карта догружается.")
        else: app.v.harta.opreste(); app.mesaj("Карта выключена.")
    b = tk.Label(d, text="OK", bg=ALBASTRU, fg="white", font=F_MIC, padx=16, pady=4); b.grid(row=9, column=0, sticky="e", padx=10, pady=10)
    b.bind("<Button-1>", ok); d.bind("<Return>", ok)


def meniu_bing(app):
    messagebox.showinfo("Карты Bing", "Microsoft закрыл бесплатный доступ к картам Bing (Bing Maps for Enterprise) для новых пользователей.\n\n"
                        "Используйте «Из карт Esri»: спутник, улицы или топография – те же функции, без ключа.", parent=app.root)
    meniu_esri(app)
