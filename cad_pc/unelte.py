"""CAD PC 3.3 – unelte adăugate după capturile din AutoCAD: tabel, hașură / gradient / contur, dispecerul de straturi,
stilurile de text, stilul punctelor, ordinea de desenare, grupuri, izolare, straturi (panoul extins), panouri extinse."""
import math, os, re
import tkinter as tk
from tkinter import ttk, messagebox, simpledialog
from ezdxf import colors as ezcolors
from stil import *

BG_P = "#3b4453"; BG_H = "#4a5a78"


def buton(p, text, cmd, **kw):
    b = tk.Label(p, text=text, bg=kw.pop("bg", FILE), fg=TEXT, font=F_MIC, padx=kw.pop("padx", 10), pady=kw.pop("pady", 4), **kw)
    b.bind("<Button-1>", lambda e: cmd()); b.bind("<Enter>", lambda e: b.config(bg=HOVER)); b.bind("<Leave>", lambda e: b.config(bg=FILE))
    return b


def dialog(app, titlu, geom=None):
    d = tk.Toplevel(app.root); d.title(titlu); d.configure(bg=CORP); d.transient(app.root)
    if geom: d.geometry(geom)
    d.bind("<Escape>", lambda e: d.destroy()); return d


def eticheta(p, text, r, c, **kw):
    l = tk.Label(p, text=text, bg=kw.pop("bg", CORP), fg=kw.pop("fg", TEXT), font=F_MIC, **kw); l.grid(row=r, column=c, sticky="w", padx=6, pady=3); return l


def intrare(p, val, r, c, w=10):
    e = tk.Entry(p, bg="#2b313a", fg=TEXT, insertbackground=TEXT, relief="flat", width=w); e.insert(0, str(val)); e.grid(row=r, column=c, sticky="w", padx=6, pady=3); return e


def numar(w, impl=0.0):
    try: return float(w.get().replace(",", "."))
    except ValueError: return impl


class Popup:
    """panou în fereastră (nu Toplevel – pe Windows acelea nu apăreau), se închide la clic în afară"""
    deschis = None

    def __init__(self, app, w, latime=None):
        if Popup.deschis: Popup.deschis.inchide()
        self.app = app; r = app.root
        self.f = tk.Frame(r, bg=BG_P, highlightthickness=1, highlightbackground="#6b778a")
        x = w.winfo_rootx() - r.winfo_rootx(); y = w.winfo_rooty() - r.winfo_rooty() + w.winfo_height()
        self.x, self.y = x, y; Popup.deschis = self
        r.after(10, lambda: r.bind_all("<Button-1>", self._afara, add="+"))

    def arata(self):
        r = self.app.root; self.f.update_idletasks()
        x = min(self.x, r.winfo_width() - self.f.winfo_reqwidth() - 4); self.f.place(x=max(0, x), y=self.y); self.f.lift()

    def _afara(self, e):
        w = e.widget
        while w is not None:
            if w is self.f: return
            w = getattr(w, "master", None)
        self.inchide()

    def inchide(self):
        try: self.app.root.unbind_all("<Button-1>")
        except Exception: pass
        try: self.f.destroy()
        except Exception: pass
        if Popup.deschis is self: Popup.deschis = None
        # legăturile globale pe care le pun listele de straturi / culori se refac la următoarea deschidere

    def rand_iconite(self, iconite):
        r = tk.Frame(self.f, bg=BG_P); r.pack(fill="x", padx=4, pady=2)
        for ic, cul, cmd, sfat in iconite:
            b = tk.Label(r, text=ic, bg=BG_P, fg=cul, font=F_ICON_MIC, width=2, padx=2)
            b.pack(side="left", padx=1); self._leaga(b, cmd); Sfat(b, sfat)
        return r

    def element(self, ic, text, cmd, cul="#e8ebef"):
        r = tk.Frame(self.f, bg=BG_P); r.pack(fill="x")
        tk.Label(r, text=ic, bg=BG_P, fg=cul, font=F_ICON_MIC, width=3).pack(side="left")
        tk.Label(r, text=text, bg=BG_P, fg=TEXT, font=F_MIC, anchor="w").pack(side="left", fill="x", expand=True, padx=(2, 12), pady=3)
        self._leaga(r, cmd); return r

    def _leaga(self, w, cmd):
        def apasa(e):
            self.inchide(); cmd() if callable(cmd) else self.app.executa(cmd)
        for x in [w] + list(w.winfo_children()):
            x.bind("<Button-1>", apasa); x.bind("<Enter>", lambda e: self._bg(w, BG_H)); x.bind("<Leave>", lambda e: self._bg(w, BG_P))

    @staticmethod
    def _bg(w, c):
        w.config(bg=c)
        for x in w.winfo_children(): x.config(bg=c)

    def separator(self): tk.Frame(self.f, bg="#5b6578", height=1).pack(fill="x", pady=2)

    def titlu(self, t): tk.Label(self.f, text="⊣  " + t, bg="#323a48", fg=GRI, font=(F, -11), anchor="w").pack(fill="x")


class Sfat:
    """sfatul de sub mouse (ca în AutoCAD): titlu, explicație, comanda. Se închide singur când mouse-ul pleacă."""
    activ = None

    def __init__(self, w, text):
        self.w = w; self.text = text; self.l = None; self.job = None
        if text:
            for x in [w] + list(w.winfo_children()):
                x.bind("<Enter>", self._intra, add="+"); x.bind("<ButtonPress>", self._iese, add="+")

    def _intra(self, e):
        if self.job: self.w.after_cancel(self.job)
        if self.l is None: self.job = self.w.after(700, self._arata)

    def _iese(self, e=None):
        if self.job:
            try: self.w.after_cancel(self.job)
            except Exception: pass
            self.job = None
        if self.l is not None:
            try: self.l.destroy()
            except Exception: pass
            self.l = None
        if Sfat.activ is self: Sfat.activ = None

    def _sub_mouse(self):
        try:
            x, y = self.w.winfo_pointerxy(); u = self.w.winfo_containing(x, y)
        except Exception: return False
        while u is not None:
            if u is self.w: return True
            u = getattr(u, "master", None)
        return False

    def _arata(self):
        self.job = None
        if not self._sub_mouse(): return
        if Sfat.activ is not None and Sfat.activ is not self: Sfat.activ._iese()       # un singur sfat deschis
        r = self.w.winfo_toplevel()
        self.l = tk.Label(r, text=self.text, bg="#f4f4f4", fg="#1b1b1b", font=(F, -12), justify="left", padx=8, pady=6, relief="solid", bd=1, wraplength=330)
        x = self.w.winfo_rootx() - r.winfo_rootx() + 10; y = self.w.winfo_rooty() - r.winfo_rooty() + self.w.winfo_height() + 4
        self.l.update_idletasks(); x = min(x, r.winfo_width() - self.l.winfo_reqwidth() - 6)
        self.l.place(x=max(0, x), y=y); self.l.lift(); Sfat.activ = self; self._verifica()

    def _verifica(self):
        if self.l is None: return
        if self._sub_mouse(): self.w.after(200, self._verifica)
        else: self._iese()


# ---------------------------------------------------------------- alegerea unei culori ACI (ca „Выбор цвета”)
def alege_aci(app, curent=7, parinte=None):
    d = tk.Toplevel(parinte or app.root); d.title("Выбор цвета"); d.configure(bg=CORP); d.transient(parinte or app.root); d.grab_set()
    rez = {"v": None}; val = tk.IntVar(value=abs(curent or 7))
    cv = tk.Canvas(d, width=24 * 12 + 4, height=22 * 20 // 2 + 60, bg=CORP, highlightthickness=0); cv.pack(padx=8, pady=8)
    k = 10
    for i in range(10, 250):
        c = "#%02x%02x%02x" % ezcolors.aci2rgb(i); x = (i - 10) // 10 * 12; y = (i - 10) % 10 * 12
        cv.create_rectangle(x + 2, y + 2, x + 13, y + 13, fill=c, outline="#20252c", tags=("c", str(i)))
    for j, i in enumerate(list(range(1, 10)) + list(range(250, 256))):
        c = "#%02x%02x%02x" % ezcolors.aci2rgb(i); cv.create_rectangle(2 + j * 18, 135, 17 + j * 18, 150, fill=c, outline="#20252c", tags=("c", str(i)))
    info = tk.Label(d, text="", bg=CORP, fg=TEXT, font=F_MIC); info.pack()
    def clic(e):
        it = cv.find_closest(e.x, e.y)
        for t in cv.gettags(it):
            if t.isdigit(): val.set(int(t)); info.config(text=f"Индекс цвета: {t}")
    cv.bind("<Button-1>", clic); cv.bind("<Double-Button-1>", lambda e: (clic(e), ok()))
    f = tk.Frame(d, bg=CORP); f.pack(fill="x", padx=8, pady=6)
    for t, v in (("ПоСлою", 256), ("ПоБлоку", 0)):
        buton(f, t, lambda v=v: (val.set(v), info.config(text=t))).pack(side="left", padx=2)
    e = tk.Entry(f, textvariable=val, width=5, bg="#2b313a", fg=TEXT, insertbackground=TEXT, relief="flat"); e.pack(side="left", padx=8)
    def ok():
        rez["v"] = val.get(); d.destroy()
    buton(f, "OK", ok, bg=ALBASTRU).pack(side="right")
    d.wait_window(); return rez["v"]


def culoare_hex(aci):
    aci = abs(aci or 7)
    return "#ffffff" if aci in (7, 255, 256, 0) else "#%02x%02x%02x" % ezcolors.aci2rgb(aci)


GROSIMI = [-3, 0, 5, 9, 13, 15, 18, 20, 25, 30, 35, 40, 50, 53, 60, 70, 80, 90, 100, 106, 120, 140, 158, 200, 211]
MODELE = ["SOLID", "ANSI31", "ANSI32", "ANSI33", "ANSI34", "ANSI35", "ANSI36", "ANSI37", "ANSI38", "AR-B816", "AR-BRSTD", "AR-CONC",
          "AR-HBONE", "AR-PARQ1", "AR-SAND", "BRICK", "CROSS", "DASH", "DOTS", "EARTH", "ESCHER", "GRASS", "GRATE", "HEX", "HONEY",
          "INSUL", "LINE", "NET", "NET3", "SQUARE", "STARS", "TRIANG", "ZIGZAG"]


def poligoane_inchise(doc, handles=None):
    """contururile închise din model: (handle, [(x, y)…], aria)"""
    from ezdxf import path as ezpath
    rez = []
    for e in doc.modelspace():
        if handles is not None and e.dxf.handle not in handles: continue
        t = e.dxftype()
        try:
            if t == "CIRCLE" or (t in ("LWPOLYLINE", "POLYLINE") and e.is_closed) or (t == "ELLIPSE" and abs(e.dxf.end_param - e.dxf.start_param - 2 * math.pi) < 1e-6) \
                    or (t == "SPLINE" and e.closed):
                pts = [(v.x, v.y) for v in ezpath.make_path(e).flattening(0.01)]
            elif t in ("LWPOLYLINE", "POLYLINE"):
                pts = [(v.x, v.y) for v in ezpath.make_path(e).flattening(0.01)]
                if len(pts) < 4 or math.hypot(pts[0][0] - pts[-1][0], pts[0][1] - pts[-1][1]) > 1e-6: continue
            else: continue
        except Exception: continue
        if len(pts) >= 3:
            a = abs(sum(pts[i][0] * pts[(i + 1) % len(pts)][1] - pts[(i + 1) % len(pts)][0] * pts[i][1] for i in range(len(pts)))) / 2
            rez.append((e.dxf.handle, pts, a))
    return rez


def in_poligon(x, y, pts):
    ins = False; n = len(pts); j = n - 1
    for i in range(n):
        xi, yi = pts[i]; xj, yj = pts[j]
        if (yi > y) != (yj > y) and x < (xj - xi) * (y - yi) / (yj - yi + 1e-300) + xi: ins = not ins
        j = i
    return ins


class Unelte:
    def __init__(self, app):
        self.app = app; self.hasura = dict(tip="Штриховка", model="ANSI31", scara=1.0, unghi=0.0, culoare=256, c1=5, c2=7, insule=True)
        self.tabel = dict(col=5, lat=25.0, rand=1, inalt=1, h=2.5, zona=False)
        self.izolare = None; self.estompare = tk.IntVar(value=50)

    @property
    def doc(self): return self.app.doc

    @property
    def ed(self): return self.app.ed

    # ================================================================ ТАБЛИЦА (Вставка таблицы)
    def tabel_dialog(self):
        if not self.doc: return self.app.mesaj("Нет открытого чертежа.")
        T = self.tabel; d = dialog(self.app, "Вставка таблицы")
        st = tk.LabelFrame(d, text="Стиль таблиц", bg=CORP, fg=TEXT, font=F_MIC); st.grid(row=0, column=0, sticky="nsew", padx=8, pady=4)
        cb = ttk.Combobox(st, values=["Standard"], state="readonly", width=18); cb.set("Standard"); cb.grid(row=0, column=0, padx=6, pady=4)
        pv = tk.LabelFrame(d, text="Параметры вставки", bg=CORP, fg=TEXT, font=F_MIC); pv.grid(row=1, column=0, sticky="nsew", padx=8, pady=4)
        sursa = tk.StringVar(value="gol")
        for i, (t, v) in enumerate((("С пустой таблицы", "gol"), ("На основе связи с данными (Excel)", "xls"), ("На основе данных объекта на чертеже (извлечение данных)", "atr"))):
            tk.Radiobutton(pv, text=t, value=v, variable=sursa, bg=CORP, fg=TEXT, selectcolor="#2b313a", activebackground=CORP, font=F_MIC,
                           command=lambda: previz()).grid(row=i, column=0, sticky="w", padx=6)
        cv = tk.Canvas(d, width=310, height=180, bg="#1f252d", highlightthickness=1, highlightbackground="#555"); cv.grid(row=2, column=0, padx=8, pady=6)
        sp = tk.LabelFrame(d, text="Способ вставки", bg=CORP, fg=TEXT, font=F_MIC); sp.grid(row=0, column=1, sticky="nsew", padx=8, pady=4)
        mod = tk.StringVar(value="zona" if T["zona"] else "punct")
        for i, (t, v) in enumerate((("Запрос точки вставки", "punct"), ("Запрос занимаемой области", "zona"))):
            tk.Radiobutton(sp, text=t, value=v, variable=mod, bg=CORP, fg=TEXT, selectcolor="#2b313a", activebackground=CORP, font=F_MIC).grid(row=i, column=0, sticky="w", padx=6)
        pc = tk.LabelFrame(d, text="Параметры столбцов и строк", bg=CORP, fg=TEXT, font=F_MIC); pc.grid(row=1, column=1, sticky="nsew", padx=8, pady=4)
        eticheta(pc, "Столбцов:", 0, 0); e_col = tk.Spinbox(pc, from_=1, to=50, width=6, command=lambda: previz()); e_col.grid(row=1, column=0, padx=6)
        eticheta(pc, "Ширина столбца:", 0, 1); e_lat = intrare(pc, T["lat"], 1, 1, 8)
        eticheta(pc, "Строк данных:", 2, 0); e_rand = tk.Spinbox(pc, from_=1, to=500, width=6, command=lambda: previz()); e_rand.grid(row=3, column=0, padx=6)
        eticheta(pc, "Высота строки (строк текста):", 2, 1); e_inalt = tk.Spinbox(pc, from_=1, to=10, width=6); e_inalt.grid(row=3, column=1, padx=6, sticky="w")
        eticheta(pc, "Высота текста:", 4, 0); e_h = intrare(pc, T["h"], 4, 1, 8)
        for w, v in ((e_col, T["col"]), (e_rand, T["rand"]), (e_inalt, T["inalt"])): w.delete(0, "end"); w.insert(0, v)
        cs = tk.LabelFrame(d, text="Задание стилей ячеек", bg=CORP, fg=TEXT, font=F_MIC); cs.grid(row=2, column=1, sticky="nsew", padx=8, pady=4)
        stiluri = []
        for i, (t, v) in enumerate((("Стиль ячеек первой строки:", "Название"), ("Стиль ячеек второй строки:", "Заголовок"), ("Стили ячеек остальных строк:", "Данные"))):
            eticheta(cs, t, i, 0); c = ttk.Combobox(cs, values=["Название", "Заголовок", "Данные"], state="readonly", width=11); c.set(v); c.grid(row=i, column=1, padx=6, pady=3)
            c.bind("<<ComboboxSelected>>", lambda e: previz()); stiluri.append(c)
        def previz():
            cv.delete("all")
            try: nc, nr = max(1, int(e_col.get())), max(1, int(e_rand.get()))
            except ValueError: return
            nc = min(nc, 8); rows = [s.get() for s in stiluri[:2]] + [stiluri[2].get()] * min(nr, 8)
            x0, y0, w, h = 16, 12, 278, 16
            for i, s in enumerate(rows):
                y = y0 + i * h
                if s == "Название":
                    cv.create_rectangle(x0, y, x0 + w, y + h, outline="white"); cv.create_text(x0 + w / 2, y + h / 2, text="Название", fill="white", font=(F, -12))
                else:
                    for j in range(nc):
                        cw = w / nc; cv.create_rectangle(x0 + j * cw, y, x0 + (j + 1) * cw, y + h, outline="white")
                        cv.create_text(x0 + j * cw + cw / 2, y + h / 2, text=s, fill="white", font=(F, -9))
        previz()
        def ok(e=None):
            try: T.update(col=max(1, int(e_col.get())), rand=max(1, int(e_rand.get())), inalt=max(1, int(e_inalt.get())), lat=numar(e_lat, 25), h=numar(e_h, 2.5))
            except ValueError: return messagebox.showerror("Таблица", "Неверное число.", parent=d)
            T["zona"] = mod.get() == "zona"; stil = [s.get() for s in stiluri]; src = sursa.get(); d.destroy()
            date = None
            if src == "xls": date = self._date_excel()
            elif src == "atr": date = self._date_atribute()
            if src != "gol" and not date: return
            if date: T["col"] = len(date[0]); T["rand"] = len(date) - 1
            if T["zona"]:
                self.ed.incepe("PUNCTE", "Первый угол таблицы:", nevoie_sel=False, n=2, prompturi=["", "Второй угол:"],
                               cb=lambda p: self._tabel(p[0], stil, date, p[1]))
            else: self.ed.incepe("PUNCT", "Точка вставки таблицы:", nevoie_sel=False, cb=lambda p: self._tabel(p, stil, date))
        f = tk.Frame(d, bg=CORP); f.grid(row=3, column=0, columnspan=2, sticky="e", padx=8, pady=8)
        buton(f, "OK", ok, bg=ALBASTRU).pack(side="left", padx=4); buton(f, "Отмена", d.destroy).pack(side="left")
        d.bind("<Return>", ok)

    def _date_excel(self):
        from tkinter import filedialog
        p = filedialog.askopenfilename(title="Связь с данными – файл Excel", filetypes=[("Excel", "*.xlsx *.xlsm")])
        if not p: return None
        from openpyxl import load_workbook
        ws = load_workbook(p, data_only=True, read_only=True).active
        rows = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(values_only=True)][:500]
        rows = [r for r in rows if any(x.strip() for x in r)]
        if not rows: self.app.mesaj("Лист Excel пуст."); return None
        n = max(len(r) for r in rows); return [r + [""] * (n - len(r)) for r in rows]

    def _date_atribute(self):
        ins = [self.ed.ent(h) for h in self.ed.sel] if self.ed.sel else list(self.doc.modelspace().query("INSERT"))
        ins = [e for e in ins if e is not None and e.dxftype() == "INSERT" and e.attribs]
        if not ins: self.app.mesaj("Нет блоков с атрибутами (выберите блоки или оставьте пустой выбор = весь чертеж)."); return None
        tags = []
        for e in ins:
            for a in e.attribs:
                if a.dxf.tag not in tags: tags.append(a.dxf.tag)
        rows = [["Блок"] + tags]
        for e in ins:
            v = {a.dxf.tag: a.dxf.text for a in e.attribs}; rows.append([e.dxf.name] + [v.get(t, "") for t in tags])
        return rows

    def _tabel(self, p, stil, date=None, p2=None):
        from ezdxf.enums import TextEntityAlignment
        T = self.tabel; h = T["h"]; nc = T["col"]; lat = T["lat"]; rh = h * (1 + T["inalt"]) * 0.9 + h * 0.6
        randuri = list(stil[:2]) + [stil[2]] * T["rand"]
        if date: randuri = [stil[1]] + [stil[2]] * (len(date) - 1)
        x0, y0 = p
        if p2 is not None:                                          # zona: lățimea și înălțimea rândurilor din dreptunghi
            x0, x1 = min(p[0], p2[0]), max(p[0], p2[0]); y1, yb = max(p[1], p2[1]), min(p[1], p2[1]); y0 = y1
            lat = (x1 - x0) / nc; rh = (y1 - yb) / len(randuri)
        msp = self.doc.modelspace(); self.ed.instantaneu([]); s = self.app.strat_curent(); W = lat * nc
        at = {"layer": s}; noi = []
        stil_txt = self.doc.header.get("$TEXTSTYLE", "Standard")
        if stil_txt not in self.doc.styles: stil_txt = "Standard"
        y = y0
        noi.append(msp.add_line((x0, y), (x0 + W, y), dxfattribs=at))
        for i, sr in enumerate(randuri):
            yb = y - rh
            noi.append(msp.add_line((x0, yb), (x0 + W, yb), dxfattribs=at)); noi.append(msp.add_line((x0, y), (x0, yb), dxfattribs=at))
            noi.append(msp.add_line((x0 + W, y), (x0 + W, yb), dxfattribs=at))
            if sr == "Название":
                t = date[i][0] if date else "Название"
                tx = msp.add_text(t, height=h * 1.3, dxfattribs=dict(at, style=stil_txt)); tx.set_placement((x0 + W / 2, y - rh / 2), align=TextEntityAlignment.MIDDLE_CENTER); noi.append(tx)
            else:
                for j in range(nc):
                    if j: noi.append(msp.add_line((x0 + j * lat, y), (x0 + j * lat, yb), dxfattribs=at))
                    t = date[i][j] if date else ""
                    if not date and sr == "Заголовок": t = f"Заголовок {j + 1}"
                    if t:
                        tx = msp.add_text(str(t), height=h * (1.0 if sr == "Данные" else 1.1), dxfattribs=dict(at, style=stil_txt))
                        tx.set_placement((x0 + j * lat + lat / 2, y - rh / 2), align=TextEntityAlignment.MIDDLE_CENTER); noi.append(tx)
            y = yb
        try:
            g = self.doc.groups.new(); g.extend(noi)                # tabelul = un grup: se selectează / mută împreună
        except Exception: pass
        self.ed.sel = {e.dxf.handle for e in noi}; self.ed.gata_modificare(f"Таблица {nc} × {len(randuri)} вставлена. Двойной щелчок по тексту – изменить ячейку.")
        self.app.marcheaza()

    # ================================================================ ШТРИХОВКА / ГРАДИЕНТ / КОНТУР
    def hasura_dialog(self, tip="Штриховка"):
        if not self.doc: return self.app.mesaj("Нет открытого чертежа.")
        H = self.hasura; H["tip"] = tip; d = dialog(self.app, "Штриховка и градиент")
        nb = tk.Frame(d, bg=CORP); nb.pack(fill="x", padx=8, pady=6)
        v_tip = tk.StringVar(value=tip)
        for t in ("Штриховка", "Градиент"):
            tk.Radiobutton(nb, text=t, value=t, variable=v_tip, bg=CORP, fg=TEXT, selectcolor="#2b313a", activebackground=CORP, font=F_MIC).pack(side="left", padx=6)
        g = tk.Frame(d, bg=CORP); g.pack(fill="both", padx=8)
        eticheta(g, "Образец:", 0, 0); cm = ttk.Combobox(g, values=MODELE, width=16); cm.set(H["model"]); cm.grid(row=0, column=1, padx=6, pady=3)
        eticheta(g, "Масштаб:", 1, 0); es = intrare(g, H["scara"], 1, 1)
        eticheta(g, "Угол:", 2, 0); eu = intrare(g, H["unghi"], 2, 1)
        culori = {"c": H["culoare"], "c1": H["c1"], "c2": H["c2"]}
        def bcul(k, r, text):
            eticheta(g, text, r, 0)
            b = tk.Label(g, text="      ", bg=culoare_hex(culori[k]), relief="solid", bd=1); b.grid(row=r, column=1, sticky="w", padx=6)
            def alege(e):
                v = alege_aci(self.app, culori[k], d)
                if v is not None: culori[k] = v; b.config(bg=culoare_hex(v))
            b.bind("<Button-1>", alege)
        bcul("c", 3, "Цвет штриховки:"); bcul("c1", 4, "Градиент – цвет 1:"); bcul("c2", 5, "Градиент – цвет 2:")
        v_ins = tk.BooleanVar(value=H["insule"])
        tk.Checkbutton(g, text="Обнаружение островков", variable=v_ins, bg=CORP, fg=TEXT, selectcolor="#2b313a", activebackground=CORP, font=F_MIC).grid(row=6, column=0, columnspan=2, sticky="w", padx=4)
        tk.Label(d, text="Контуры: выбранные замкнутые объекты, иначе – укажите точки внутри областей (Esc – конец).",
                 bg=CORP, fg=GRI, font=(F, -11), wraplength=360, justify="left").pack(padx=8, pady=4, anchor="w")
        def ok(e=None):
            H.update(tip=v_tip.get(), model=cm.get().strip().upper() or "ANSI31", scara=numar(es, 1) or 1, unghi=numar(eu, 0),
                     culoare=culori["c"], c1=culori["c1"], c2=culori["c2"], insule=v_ins.get()); d.destroy(); self.hasura_porneste()
        f = tk.Frame(d, bg=CORP); f.pack(fill="x", padx=8, pady=8)
        buton(f, "OK", ok, bg=ALBASTRU).pack(side="right"); buton(f, "Отмена", d.destroy).pack(side="right", padx=4)
        d.bind("<Return>", ok)

    def hasura_porneste(self):
        cont = poligoane_inchise(self.doc, set(self.ed.sel)) if self.ed.sel else []
        if cont:
            self.ed.instantaneu([]); n = 0
            for _, pts, _ in cont: self._hasura(pts, []); n += 1
            self.ed.gata_modificare(f"{self.hasura['tip']}: {n} контур(ов)."); return self.app.marcheaza()
        self._cont = None
        self.ed.incepe("REPETAT", "Укажите внутреннюю точку (Esc – завершить):", nevoie_sel=False, rep="Укажите внутреннюю точку (Esc – завершить):",
                       cb=self._hasura_punct)

    def _gaseste_contur(self, p):
        if self._cont is None: self._cont = poligoane_inchise(self.doc)
        cand = sorted([c for c in self._cont if in_poligon(p[0], p[1], c[1])], key=lambda c: c[2])
        if not cand: return None, []
        ext = cand[0]; insule = []
        if self.hasura["insule"]:                                  # contururile închise din interior = insule (rămân goale)
            for c in self._cont:
                if c is ext or c[2] >= ext[2]: continue
                if all(in_poligon(x, y, ext[1]) for x, y in c[1][::max(1, len(c[1]) // 8)]):
                    if not any(c is not o and c[2] < o[2] and in_poligon(c[1][0][0], c[1][0][1], o[1]) for o in insule): insule.append(c)
        return ext, insule

    def _hasura_punct(self, p):
        ext, insule = self._gaseste_contur(p)
        if ext is None: return self.app.mesaj("Замкнутый контур вокруг точки не найден (нужны замкнутые полилинии, круги, эллипсы).")
        self.ed.instantaneu([]); self._hasura(ext[1], [i[1] for i in insule]); self.ed.gata_modificare(f"{self.hasura['tip']} добавлена."); self.app.marcheaza()

    def _hasura(self, pts, insule):
        H = self.hasura; msp = self.doc.modelspace()
        h = msp.add_hatch(color=H["culoare"] if H["tip"] == "Штриховка" else 256, dxfattribs={"layer": self.app.strat_curent()})
        if H["tip"] == "Градиент":
            h.set_gradient(color1=ezcolors.aci2rgb(abs(H["c1"]) or 7), color2=ezcolors.aci2rgb(abs(H["c2"]) or 7), rotation=H["unghi"], name="LINEAR")
        elif H["model"] == "SOLID": h.set_solid_fill(color=H["culoare"])
        else:
            try: h.set_pattern_fill(H["model"], color=H["culoare"], angle=H["unghi"], scale=H["scara"])
            except Exception: h.set_pattern_fill("ANSI31", color=H["culoare"], angle=H["unghi"], scale=H["scara"])
        h.paths.add_polyline_path(pts, is_closed=True, flags=1)
        for i in insule: h.paths.add_polyline_path(i, is_closed=True, flags=16)
        return h

    def contur(self):
        if not self.doc: return
        self._cont = None
        def cb(p):
            ext, insule = self._gaseste_contur(p)
            if ext is None: return self.app.mesaj("Замкнутый контур вокруг точки не найден.")
            self.ed.instantaneu([]); msp = self.doc.modelspace(); s = self.app.strat_curent()
            for c in [ext] + insule: msp.add_lwpolyline(c[1], close=True, dxfattribs={"layer": s})
            self.ed.gata_modificare(f"Контур: создано полилиний {1 + len(insule)} (площадь {ext[2]:.3f})."); self.app.marcheaza()
        self.ed.incepe("REPETAT", "КОНТУР  Укажите внутреннюю точку (Esc – завершить):", nevoie_sel=False, rep="Укажите внутреннюю точку (Esc – завершить):", cb=cb)

    # ================================================================ ДИСПЕТЧЕР СВОЙСТВ СЛОЕВ
    def straturi(self):
        app = self.app
        if not self.doc: return app.mesaj("Нет открытого чертежа.")
        d = dialog(app, "Диспетчер свойств слоев", "1080x520"); d.configure(bg=BG_P)
        sus = tk.Frame(d, bg=BG_P); sus.pack(fill="x", padx=6, pady=(6, 2))
        lbl_cur = tk.Label(sus, text="", bg=BG_P, fg=TEXT, font=F_MIC); lbl_cur.pack(side="left")
        cauta = tk.Entry(sus, bg="#2b313a", fg=TEXT, insertbackground=TEXT, relief="flat", width=24); cauta.pack(side="right", padx=4)
        tk.Label(sus, text="Поиск слоя 🔎", bg=BG_P, fg=GRI, font=F_MIC).pack(side="right")
        bara = tk.Frame(d, bg=BG_P); bara.pack(fill="x", padx=6)
        corp = tk.Frame(d, bg=BG_P); corp.pack(fill="both", expand=True, padx=6, pady=4)
        filt = ttk.Treeview(corp, show="tree", style="Cad.Treeview", height=6); filt.pack(side="left", fill="y")
        filt.column("#0", width=170); filt.insert("", "end", "toate", text="Фильтры", open=True)
        filt.insert("toate", "end", "all", text="Все"); filt.insert("all", "end", "used", text="Все используемые слои"); filt.item("all", open=True)
        filt.selection_set("all")
        inv = tk.BooleanVar(value=False)
        cols = ("st", "nume", "on", "frz", "lock", "plot", "cul", "tip", "ves", "tr", "desc")
        tv = ttk.Treeview(corp, columns=cols, show="headings", style="Cad.Treeview", selectmode="extended")
        for c, t, w in zip(cols, ("С", "Имя", "Вкл", "Заморозить", "Блокировать", "Печать", "Цвет", "Тип линий", "Вес линий", "Прозрачность", "Пояснение"),
                           (30, 200, 40, 75, 80, 55, 90, 100, 85, 85, 160)):
            tv.heading(c, text=t); tv.column(c, width=w, anchor="w" if c in ("nume", "cul", "tip", "ves", "desc") else "center")
        sb = ttk.Scrollbar(corp, orient="vertical", command=tv.yview); tv.configure(yscrollcommand=sb.set)
        tv.pack(side="left", fill="both", expand=True); sb.pack(side="left", fill="y")
        jos = tk.Frame(d, bg=BG_P); jos.pack(fill="x", padx=6, pady=(0, 6))
        tk.Checkbutton(jos, text="Инвертировать фильтр", variable=inv, command=lambda: umple(), bg=BG_P, fg=TEXT, selectcolor="#2b313a", activebackground=BG_P, font=F_MIC).pack(side="left")
        stare = tk.Label(jos, text="", bg=BG_P, fg=TEXT, font=F_MIC); stare.pack(side="left", padx=20)

        def folosite():
            return {e.dxf.get("layer", "0") for e in self.doc.modelspace()} | {e.dxf.get("layer", "0") for l in self.doc.layouts for e in l}

        def umple():
            sel = tv.selection(); tv.delete(*tv.get_children()); cur = self.doc.header.get("$CLAYER", "0")
            q = cauta.get().strip().lower(); fl = filt.selection(); us = folosite() if fl and fl[0] == "used" else None
            toate = sorted(self.doc.layers, key=lambda l: l.dxf.name.lower()); n = 0
            for l in toate:
                nm = l.dxf.name; ok = (not q or q in nm.lower()) and (us is None or nm in us)
                if inv.get(): ok = not ok
                if not ok: continue
                c = l.dxf.get("color", 7); aci = abs(c)
                cul = ("■ " + ("белый" if aci == 7 else str(aci))) if not l.dxf.hasattr("true_color") else "■ " + "%d,%d,%d" % ezcolors.int2rgb(l.dxf.true_color)
                lw = l.dxf.get("lineweight", -3); ves = "—— По умолч." if lw < 0 else f"—— {lw / 100:.2f} мм"
                try: tr = int(round((l.transparency or 0) * 100))
                except Exception: tr = 0
                try: ds = l.description
                except Exception: ds = ""
                tv.insert("", "end", iid=nm, values=("✔" if nm == cur else "▱", nm, "💡" if l.is_on() else "○", "❄" if l.is_frozen() else "☀",
                                                     "🔒" if l.is_locked() else "🔓", "🖶" if l.dxf.get("plot", 1) else "⊘", cul,
                                                     l.dxf.get("linetype", "Continuous"), ves, tr, ds), tags=(nm,))
                n += 1
            for s in sel:
                if tv.exists(s): tv.selection_add(s)
            lbl_cur.config(text=f"Текущий слой: {cur}")
            stare.config(text=f"{'Все используемые' if us is not None else 'Все'}: отображается слоев — {n}, всего слоев — {len(toate)}")

        def gata():
            app.modificat = True; app._titlu(); umple(); app.v.reconstruieste(); app.v.deseneaza(); app._actualizeaza_straturi()

        def clic(e):
            col = tv.identify_column(e.x); it = tv.identify_row(e.y)
            if not it or tv.identify_region(e.x, e.y) != "cell": return
            l = self.doc.layers.get(it); cur = self.doc.header.get("$CLAYER", "0")
            if col == "#1": self.doc.header["$CLAYER"] = it
            elif col == "#3": l.off() if l.is_on() else l.on()
            elif col == "#4":
                if it == cur: return app.mesaj("Текущий слой нельзя заморозить.")
                l.thaw() if l.is_frozen() else l.freeze()
            elif col == "#5": l.unlock() if l.is_locked() else l.lock()
            elif col == "#6": l.dxf.plot = 0 if l.dxf.get("plot", 1) else 1
            elif col == "#7":
                v = alege_aci(app, l.dxf.get("color", 7), d)
                if v is None or v in (0, 256): return
                if l.dxf.hasattr("true_color"): l.dxf.discard("true_color")
                l.color = v
            elif col == "#8":
                m = tk.Menu(d, tearoff=0)
                for lt in sorted(x.dxf.name for x in self.doc.linetypes):
                    m.add_command(label=lt, command=lambda lt=lt: (setattr(l.dxf, "linetype", lt), gata()))
                return m.tk_popup(e.x_root, e.y_root)
            elif col == "#9":
                m = tk.Menu(d, tearoff=0)
                for w in GROSIMI: m.add_command(label="По умолчанию" if w < 0 else f"{w / 100:.2f} мм", command=lambda w=w: (setattr(l.dxf, "lineweight", w), gata()))
                return m.tk_popup(e.x_root, e.y_root)
            elif col == "#10":
                v = simpledialog.askinteger("Прозрачность", "Прозрачность слоя (0–90):", parent=d, minvalue=0, maxvalue=90)
                if v is None: return
                l.transparency = v / 100
            else: return
            gata()

        def dublu(e):
            col = tv.identify_column(e.x); it = tv.identify_row(e.y)
            if not it: return
            l = self.doc.layers.get(it)
            if col == "#2":
                if it in ("0", "Defpoints"): return app.mesaj("Слои 0 и Defpoints нельзя переименовать.")
                n = simpledialog.askstring("Переименовать слой", "Новое имя:", initialvalue=it, parent=d)
                if n and n != it and n not in self.doc.layers:
                    try:
                        l.rename(n)
                        if self.doc.header.get("$CLAYER") == it: self.doc.header["$CLAYER"] = n
                    except Exception as ex: return messagebox.showerror("Слой", str(ex), parent=d)
                    gata()
            elif col == "#11":
                n = simpledialog.askstring("Пояснение", "Пояснение слоя:", initialvalue=getattr(l, "description", ""), parent=d)
                if n is not None:
                    try: l.description = n
                    except Exception: pass
                    gata()
            elif col == "#1": self.doc.header["$CLAYER"] = it; gata()
        tv.bind("<Button-1>", lambda e: d.after(1, lambda: clic(e))); tv.bind("<Double-Button-1>", dublu)
        cauta.bind("<KeyRelease>", lambda e: umple()); filt.bind("<<TreeviewSelect>>", lambda e: umple())

        def nou(inghetat=False):
            i = 1
            while f"Слой{i}" in self.doc.layers: i += 1
            l = self.doc.layers.add(f"Слой{i}", color=7)
            if inghetat: l.freeze()
            gata(); tv.selection_set(f"Слой{i}"); tv.see(f"Слой{i}")
        def sterge():
            cur = self.doc.header.get("$CLAYER", "0"); folos = folosite(); n = 0
            for it in tv.selection():
                if it in ("0", "Defpoints") or it == cur: app.mesaj(f"Слой «{it}» нельзя удалить (0, Defpoints или текущий)."); continue
                if it in folos: app.mesaj(f"Слой «{it}» содержит объекты – не удален (используйте «Удалить слой» в панели «Слои»)."); continue
                self.doc.layers.remove(it); n += 1
            gata(); app.mesaj(f"Удалено слоев: {n}.")
        def curent():
            s = tv.selection()
            if s: self.doc.header["$CLAYER"] = s[0]; gata()
        for ic, sfat, f in (("✚", "Создать слой", nou), ("❄✚", "Новый слой, замороженный на всех ВЭ", lambda: nou(True)), ("✖", "Удалить слой", sterge),
                            ("✔", "Установить текущий", curent), ("⟳", "Обновить", umple)):
            b = tk.Label(bara, text=ic, bg=BG_P, fg="#f0c040" if ic != "✖" else "#e05050", font=F_ICON_MIC, padx=6); b.pack(side="left")
            b.bind("<Button-1>", lambda e, f=f: f()); Sfat(b, sfat)
        tk.Label(bara, text="Щелчок по значку – изменить;  двойной щелчок по имени – переименовать", bg=BG_P, fg=GRI, font=(F, -11)).pack(side="right")
        umple()

    # ================================================================ panoul „Слои” extins
    def _straturi_sel(self):
        s = {self.ed.ent(h).dxf.get("layer", "0") for h in self.ed.sel if self.ed.ent(h) is not None}
        if not s: self.app.mesaj("Выберите объекты на нужных слоях, затем команду.")
        return s

    def _gata_straturi(self, msg):
        self.app.modificat = True; self.app._titlu(); self.app.v.reconstruieste(); self.app.v.deseneaza(); self.app._actualizeaza_straturi(); self.app.mesaj(msg)

    def strat_op(self, op):
        doc = self.doc
        if not doc: return
        L = doc.layers; cur = doc.header.get("$CLAYER", "0")
        if op in ("ВКЛ", "РАЗМ", "РАЗБЛОК"):
            for l in L: (l.on() if op == "ВКЛ" else l.thaw() if op == "РАЗМ" else l.unlock())
            return self._gata_straturi({"ВКЛ": "Все слои включены.", "РАЗМ": "Все слои разморожены.", "РАЗБЛОК": "Все слои разблокированы."}[op])
        if op == "НЕИЗО":
            if not self.izolare: return self.app.mesaj("Нет изолированных слоев.")
            for n, (on, fr, lk) in self.izolare.items():
                if n in L:
                    l = L.get(n); l.on() if on else l.off(); (l.freeze() if fr else l.thaw()) if n != cur else None; l.lock() if lk else l.unlock()
            self.izolare = None; return self._gata_straturi("Изоляция слоев отменена.")
        s = self._straturi_sel()
        if not s: return
        if op == "ИЗО":
            self.izolare = {l.dxf.name: (l.is_on(), l.is_frozen(), l.is_locked()) for l in L}
            for l in L:
                if l.dxf.name not in s: l.off()
            if cur not in s: doc.header["$CLAYER"] = sorted(s)[0]
            return self._gata_straturi(f"Изолированы слои: {', '.join(sorted(s))}. «Отменить изоляцию слоев» – вернуть.")
        if op == "ТЕК":
            doc.header["$CLAYER"] = sorted(s)[0]; return self._gata_straturi(f"Текущий слой: {sorted(s)[0]}")
        if op in ("ОТКЛ", "ЗАМ", "БЛОК"):
            for n in s:
                if n == cur and op != "БЛОК": continue
                l = L.get(n); l.off() if op == "ОТКЛ" else l.freeze() if op == "ЗАМ" else l.lock()
            self.ed.goleste(); return self._gata_straturi({"ОТКЛ": "Слои отключены: ", "ЗАМ": "Слои заморожены: ", "БЛОК": "Слои блокированы: "}[op] + ", ".join(sorted(s)))
        if op == "НАТЕК":
            return self.ed.schimba_strat(cur)
        if op == "ОБЪЕД":
            m = simpledialog.askstring("Объединить слои", f"Объекты слоев {', '.join(sorted(s))} перенести на слой:", initialvalue=cur, parent=self.app.root)
            if not m: return
            if m not in L: L.add(m)
            hs = [e.dxf.handle for e in doc.modelspace() if e.dxf.get("layer", "0") in s]; self.ed.instantaneu(hs, complet=True)
            for e in doc.modelspace():
                if e.dxf.get("layer", "0") in s: e.dxf.layer = m
            for n in s:
                if n not in ("0", "Defpoints", m) and n != cur:
                    try: L.remove(n)
                    except Exception: pass
            self.ed.goleste(); return self._gata_straturi(f"Слои объединены в «{m}».")
        if op == "УДАЛ":
            s = {n for n in s if n not in ("0", "Defpoints", cur)}
            if not s: return self.app.mesaj("Слои 0, Defpoints и текущий не удаляются.")
            if not messagebox.askyesno("Удалить слой", f"Удалить слои {', '.join(sorted(s))} и ВСЕ объекты на них?", parent=self.app.root): return
            self.ed.instantaneu(None, complet=True); msp = doc.modelspace()
            for e in [e for e in msp if e.dxf.get("layer", "0") in s]: msp.delete_entity(e)
            for n in s: L.remove(n)
            self.ed.goleste(); return self._gata_straturi(f"Удалены слои и их объекты: {', '.join(sorted(s))}.")

    def estompare_blocate(self, v=None):
        if v is not None: self.estompare.set(v)
        self.app.v.estompare = self.estompare.get(); self.app.v.reconstruieste(); self.app.v.deseneaza()
        self.app.mesaj(f"Затенение блокированных слоев: {self.estompare.get()}%")

    # ================================================================ stiluri de text
    def lista_stiluri(self, w):
        if not self.doc: return
        p = Popup(self.app, w); cur = self.doc.header.get("$TEXTSTYLE", "Standard")
        rand = tk.Frame(p.f, bg=BG_P); rand.pack(padx=6, pady=6)
        for i, s in enumerate(sorted(x.dxf.name for x in self.doc.styles if x.dxf.name)):
            st = self.doc.styles.get(s); font = (st.dxf.get("font", "") or "").lower()
            fam = "Arial" if "arial" in font else "Times New Roman" if "times" in font else "Consolas" if "cour" in font else F
            c = tk.Frame(rand, bg=BG_H if s == cur else BG_P, padx=3, pady=3); c.grid(row=i // 3, column=i % 3, padx=3, pady=3)
            tk.Label(c, text="Ao Bb12", bg="white", fg="#111", font=(fam, -18), width=9, height=2).pack()
            tk.Label(c, text=s, bg=c["bg"], fg=TEXT, font=F_MIC).pack()
            p._leaga(c, lambda s=s: self.seteaza_stil(s))
        p.separator(); p.element("", "Управление стилями текста...", self.stiluri_dialog); p.arata()

    def seteaza_stil(self, s):
        self.doc.header["$TEXTSTYLE"] = s; cb = self.app.pg.combo.get("стильтекста")
        if cb is not None:
            try: cb.set(s)
            except Exception: pass
        if self.ed.sel:
            hs = [h for h in self.ed.sel if self.ed.ent(h) is not None and self.ed.ent(h).dxftype() in ("TEXT", "MTEXT")]
            if hs:
                self.ed.instantaneu(hs)
                for h in hs: self.ed.ent(h).dxf.style = s
                return self.ed.gata_modificare(f"Стиль «{s}» применен к {len(hs)} текстам.")
        self.app.mesaj(f"Текущий текстовый стиль: {s}")

    def stiluri_dialog(self):
        if not self.doc: return
        d = dialog(self.app, "Текстовые стили"); doc = self.doc
        lb = tk.Listbox(d, bg="#2b313a", fg=TEXT, selectbackground=ACTIV, height=12, width=24, exportselection=False); lb.grid(row=0, column=0, rowspan=8, padx=8, pady=8, sticky="ns")
        eticheta(d, "Шрифт:", 0, 1)
        cf = ttk.Combobox(d, values=["txt.shx", "simplex.shx", "romans.shx", "isocp.shx", "gost type a.shx", "arial.ttf", "arialn.ttf", "times.ttf", "calibri.ttf",
                                     "cour.ttf", "isocpeur.ttf", "GOST type A.ttf"], width=22); cf.grid(row=1, column=1, padx=6, sticky="w")
        eticheta(d, "Высота (0 = задается при вводе):", 2, 1); eh = intrare(d, 0, 3, 1)
        eticheta(d, "Степень растяжения:", 4, 1); ew = intrare(d, 1, 5, 1)
        eticheta(d, "Угол наклона:", 6, 1); eo = intrare(d, 0, 7, 1)
        prev = tk.Label(d, text="AaBb123 АаБб", bg="white", fg="#111", font=(F, -22), width=16); prev.grid(row=8, column=0, columnspan=2, pady=6)
        def umple():
            lb.delete(0, "end")
            for s in sorted(x.dxf.name for x in doc.styles if x.dxf.name): lb.insert("end", s)
        def ales(e=None):
            s = lb.get(lb.curselection()[0]) if lb.curselection() else None
            if not s: return
            st = doc.styles.get(s)
            for w, v in ((eh, st.dxf.get("height", 0)), (ew, st.dxf.get("width", 1)), (eo, st.dxf.get("oblique", 0))): w.delete(0, "end"); w.insert(0, v)
            cf.set(st.dxf.get("font", "txt.shx"))
            f = (st.dxf.get("font", "") or "").lower(); prev.config(font=("Arial" if "arial" in f else "Times New Roman" if "times" in f else F, -22))
        lb.bind("<<ListboxSelect>>", ales)
        def aplica():
            s = lb.get(lb.curselection()[0]) if lb.curselection() else None
            if not s: return
            st = doc.styles.get(s); st.dxf.font = cf.get() or "txt.shx"; st.dxf.height = numar(eh, 0); st.dxf.width = numar(ew, 1) or 1; st.dxf.oblique = numar(eo, 0)
            self.app.marcheaza(); self.app.v.reconstruieste(); self.app.v.deseneaza(); self.app.mesaj(f"Стиль «{s}» изменен.")
        def nou():
            n = simpledialog.askstring("Новый текстовый стиль", "Имя стиля:", initialvalue="Стиль1", parent=d)
            if n and n not in doc.styles: doc.styles.add(n, font="arial.ttf"); umple(); self.app.marcheaza()
        def sterge():
            s = lb.get(lb.curselection()[0]) if lb.curselection() else None
            if not s or s == "Standard": return
            if any(e.dxf.get("style", "Standard") == s for e in doc.modelspace() if e.dxftype() in ("TEXT", "MTEXT")): return messagebox.showinfo("Стиль", "Стиль используется – не удален.", parent=d)
            doc.styles.remove(s); umple()
        f = tk.Frame(d, bg=CORP); f.grid(row=9, column=0, columnspan=2, sticky="e", padx=8, pady=8)
        for t, c in (("Установить текущим", lambda: lb.curselection() and self.seteaza_stil(lb.get(lb.curselection()[0]))), ("Новый...", nou), ("Удалить", sterge),
                     ("Применить", aplica), ("Закрыть", d.destroy)):
            buton(f, t, c, bg=ALBASTRU if t == "Применить" else FILE).pack(side="left", padx=2)
        umple()

    # ================================================================ stilul punctelor (Отображение точек)
    def puncte_dialog(self):
        if not self.doc: return
        d = dialog(self.app, "Отображение точек"); moduri = [0, 1, 2, 3, 4, 32, 33, 34, 35, 36, 64, 65, 66, 67, 68, 96, 97, 98, 99, 100]
        v = tk.IntVar(value=self.doc.header.get("$PDMODE", 0)); cv = tk.Canvas(d, width=5 * 54, height=4 * 54, bg=CORP, highlightthickness=0); cv.pack(padx=8, pady=8)
        def desen():
            cv.delete("all")
            for i, m in enumerate(moduri):
                x, y = i % 5 * 54 + 27, i // 5 * 54 + 27; sel = m == v.get()
                cv.create_rectangle(x - 24, y - 24, x + 24, y + 24, fill="#4a6fa5" if sel else "#2b313a", outline="#666", tags=str(m))
                b = m & 31; c = "#e8ebef"
                if b == 0: cv.create_oval(x - 1, y - 1, x + 1, y + 1, fill=c, outline=c)
                if b == 2: cv.create_line(x - 9, y, x + 9, y, fill=c); cv.create_line(x, y - 9, x, y + 9, fill=c)
                if b == 3: cv.create_line(x - 7, y - 7, x + 7, y + 7, fill=c); cv.create_line(x - 7, y + 7, x + 7, y - 7, fill=c)
                if b == 4: cv.create_line(x, y, x, y - 9, fill=c)
                if m & 32: cv.create_oval(x - 8, y - 8, x + 8, y + 8, outline=c)
                if m & 64: cv.create_rectangle(x - 8, y - 8, x + 8, y + 8, outline=c)
        def clic(e):
            for t in cv.gettags(cv.find_closest(e.x, e.y)):
                if t.isdigit(): v.set(int(t)); desen()
        cv.bind("<Button-1>", clic); desen()
        f = tk.Frame(d, bg=CORP); f.pack(fill="x", padx=8)
        eticheta(f, "Размер точки:", 0, 0); es = intrare(f, self.doc.header.get("$PDSIZE", 0.0), 0, 1)
        def ok():
            self.doc.header["$PDMODE"] = v.get(); self.doc.header["$PDSIZE"] = numar(es, 0); d.destroy(); self.app.marcheaza()
            self.app.v.reconstruieste(); self.app.v.deseneaza(); self.app.mesaj(f"Стиль точек: {v.get()}, размер {numar(es, 0)}.")
        buton(d, "OK", ok, bg=ALBASTRU).pack(side="right", padx=8, pady=8)

    # ================================================================ ordinea de desenare
    def ordine(self, op):
        if not self.doc: return
        msp = self.doc.modelspace()
        try: ordine = [e.dxf.handle for e in msp.entities_in_redraw_order()]
        except Exception: ordine = [e.dxf.handle for e in msp]
        tipuri = {"ТЕКСТВПЕРЕД": ("TEXT", "MTEXT"), "РАЗМВПЕРЕД": ("DIMENSION",), "ВЫНВПЕРЕД": ("LEADER", "MULTILEADER", "MLEADER"),
                  "АННВПЕРЕД": ("TEXT", "MTEXT", "DIMENSION", "LEADER", "MULTILEADER", "MLEADER"), "ШТРИХНАЗАД": ("HATCH",)}
        if op in tipuri:
            sel = [h for h in ordine if self.ed.ent(h) is not None and self.ed.ent(h).dxftype() in tipuri[op]]
            fata = op != "ШТРИХНАЗАД"
        else:
            sel = [h for h in ordine if h in self.ed.sel]
            if not sel: return self.app.mesaj("Выберите объекты, затем «Порядок прорисовки».")
            if op in ("ПЕРЕД", "ЗА"):
                def cb(p):
                    sx, sy = self.app.v.ecran(*p); ref = self.ed.la_entitate(sx, sy)
                    if not ref or ref in sel: return self.app.mesaj("Опорный объект не найден.")
                    rest = [h for h in ordine if h not in sel]; i = rest.index(ref)
                    nou = rest[:i + 1] + sel + rest[i + 1:] if op == "ПЕРЕД" else rest[:i] + sel + rest[i:]
                    self._aplica_ordine(msp, nou)
                return self.ed.incepe("PUNCT", "Выберите опорный объект:", nevoie_sel=False, cb=cb)
            fata = op == "ВПЕРЕД"
        rest = [h for h in ordine if h not in set(sel)]
        self._aplica_ordine(msp, rest + sel if fata else sel + rest)

    def _aplica_ordine(self, msp, nou):
        try: msp.set_redraw_order({h: "%X" % (i + 1) for i, h in enumerate(nou)})
        except Exception as ex: return self.app.mesaj(f"Порядок прорисовки: {ex}")
        self.app.marcheaza(); self.app.v.reconstruieste(); self.app.v.deseneaza(); self.app.mesaj("Порядок прорисовки изменен.")

    # ================================================================ grupuri
    def grupuri_ale(self, h):
        try: return [g for _, g in self.doc.groups if h in {e.dxf.handle for e in g}]
        except Exception: return []

    def extinde_grup(self, h):
        """la clic pe un obiect dintr-un grup se selectează tot grupul (ca în AutoCAD, cu „Выбор группы” activ)"""
        if not getattr(self, "selectie_grup", True) or not self.doc: return {h}
        hs = {h}
        for g in self.grupuri_ale(h): hs |= {e.dxf.handle for e in g if e.is_alive}
        return hs

    def grup_op(self, op):
        doc = self.doc
        if not doc: return
        if op == "ВЫБОР":
            self.selectie_grup = not getattr(self, "selectie_grup", True)
            return self.app.mesaj("Выбор группы: " + ("включен" if self.selectie_grup else "выключен"))
        if op == "ДИСП": return self.dispecer_grupuri()
        if op == "РАМКА": return self.rama_grup()
        if not self.ed.sel: return self.app.mesaj("Выберите объекты, затем команду.")
        ents = [self.ed.ent(h) for h in self.ed.sel if self.ed.ent(h) is not None]
        if op == "ГРУППА":
            n = simpledialog.askstring("Группа", "Имя группы (пусто = безымянная):", parent=self.app.root)
            if n is None: return
            try: g = doc.groups.new(n.strip() or None)
            except Exception as ex: return messagebox.showerror("Группа", str(ex))
            g.extend(ents); self.app.marcheaza(); return self.app.mesaj(f"Группа «{g.dxf.name if hasattr(g.dxf, 'name') else n}» из {len(ents)} объектов.")
        if op == "РАЗГРУППИРОВАТЬ":
            gs = set()
            for e in ents:
                for g in self.grupuri_ale(e.dxf.handle): gs.add(g)
            for nume, g in list(doc.groups):
                if g in gs: doc.groups.delete(nume)
            self.app.marcheaza(); return self.app.mesaj(f"Разгруппировано групп: {len(gs)}.")
        if op == "РЕДГРУППА":
            gs = self.grupuri_ale(ents[0].dxf.handle)
            if not gs: return self.app.mesaj("Объект не входит в группу.")
            g = gs[0]; noi = [e for e in ents if e.dxf.handle not in {x.dxf.handle for x in g}]
            if noi: g.extend(noi); return self.app.mesaj(f"В группу добавлено {len(noi)} объектов.")
            return self.app.mesaj("Выберите объекты группы и новые объекты, чтобы добавить их.")

    def rama_grup(self):
        v = self.app.v; cv = self.app.cv; cv.delete("ramagrup")
        for nume, g in self.doc.groups:
            bb = [v.bbox.get(e.dxf.handle) for e in g if e.is_alive and v.bbox.get(e.dxf.handle)]
            if not bb or not any(e.dxf.handle in self.ed.sel for e in g): continue
            x0, y0 = min(b[0] for b in bb), min(b[1] for b in bb); x1, y1 = max(b[2] for b in bb), max(b[3] for b in bb)
            a, b = v.ecran(x0, y1); c, dd = v.ecran(x1, y0)
            cv.create_rectangle(a - 4, b - 4, c + 4, dd + 4, outline="#5fb0ff", dash=(5, 3), width=2, tags=("ramagrup", "ov"))
        self.app.mesaj("Ограничивающая рамка группы показана (исчезнет при перерисовке).")

    def dispecer_grupuri(self):
        if not self.doc: return
        d = dialog(self.app, "Диспетчер групп"); tv = ttk.Treeview(d, columns=("n", "c"), show="headings", style="Cad.Treeview", height=12)
        tv.heading("n", text="Имя группы"); tv.heading("c", text="Объектов"); tv.column("c", width=80, anchor="center"); tv.pack(fill="both", expand=True, padx=8, pady=8)
        def umple():
            tv.delete(*tv.get_children())
            for nume, g in self.doc.groups: tv.insert("", "end", iid=nume, values=(nume, len(g)))
        def selecteaza():
            for it in tv.selection():
                g = self.doc.groups.get(it); self.ed.sel |= {e.dxf.handle for e in g if e.is_alive}
            self.ed.arata_selectia(); self.app.dupa_selectie()
        def sterge():
            for it in tv.selection(): self.doc.groups.delete(it)
            umple(); self.app.marcheaza()
        f = tk.Frame(d, bg=CORP); f.pack(fill="x", padx=8, pady=6)
        for t, c in (("Выделить", selecteaza), ("Разгруппировать", sterge), ("Закрыть", d.destroy)): buton(f, t, c).pack(side="left", padx=2)
        umple()

    # ================================================================ izolarea obiectelor
    def izoleaza(self, op):
        v = self.app.v
        if op == "КОНЕЦ": v.izolate = None; v.ascunse_h = set(); v.deseneaza(); return self.app.mesaj("Изоляция объектов завершена.")
        if not self.ed.sel: return self.app.mesaj("Выберите объекты.")
        if op == "ИЗОЛ": v.izolate = set(self.ed.sel); v.ascunse_h = set()
        else: v.ascunse_h = set(getattr(v, "ascunse_h", set())) | set(self.ed.sel)
        self.ed.goleste(); v.deseneaza(); self.app.mesaj("Объекты " + ("изолированы" if op == "ИЗОЛ" else "скрыты") + ". «Завершить изоляцию объектов» – вернуть все.")
