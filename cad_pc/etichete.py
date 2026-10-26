"""Etapa 2 – etichetele: text nou (ТЕКСТ / МТЕКСТ), editarea textului (dublu clic), editorul de atribute,
„Найти и заменить”, numerotarea automată (Auto Number) și exportul / importul atributelor în Excel."""
import os, re
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from stil import *

CULORI = [("ПоСлою", 256), ("ПоБлоку", 0), ("Красный", 1), ("Желтый", 2), ("Зеленый", 3), ("Голубой", 4), ("Синий", 5), ("Фиолетовый", 6), ("Белый", 7)]


def _dialog(app, titlu, w=460, h=360):
    d = tk.Toplevel(app.root); d.title(titlu); d.configure(bg=CORP); d.transient(app.root); d.geometry(f"{w}x{h}")
    d.grab_set(); return d


def _rand(d, eticheta, widget_fn, r):
    tk.Label(d, text=eticheta, bg=CORP, fg=TEXT, font=F_MIC, anchor="w").grid(row=r, column=0, sticky="w", padx=10, pady=4)
    w = widget_fn(d); w.grid(row=r, column=1, sticky="we", padx=10, pady=4); return w


def _intrare(val=""):
    def f(p):
        e = tk.Entry(p, bg="#2b313a", fg=TEXT, insertbackground=TEXT, relief="flat", font=F_MIC); e.insert(0, str(val)); return e
    return f


def _butoane(d, ok, r, text_ok="OK"):
    f = tk.Frame(d, bg=CORP); f.grid(row=r, column=0, columnspan=2, sticky="e", padx=10, pady=10)
    for t, cmd in ((text_ok, ok), ("Отмена", d.destroy)):
        b = tk.Label(f, text=t, bg=ALBASTRU if t == text_ok else FILE, fg="white", font=F_MIC, padx=16, pady=4); b.pack(side="left", padx=4)
        b.bind("<Button-1>", lambda e, c=cmd: c())
    d.bind("<Return>", lambda e: ok() if not isinstance(e.widget, tk.Text) else None); d.bind("<Escape>", lambda e: d.destroy())


class Etichete:
    def __init__(self, app):
        self.app = app; self.inaltime = 2.5

    @property
    def doc(self): return self.app.doc

    # ---------------------------------------------------------------- text nou / editare text
    def dialog_text(self, e=None, punct=None, mtext=False):
        nou = e is None
        if not nou: mtext = e.dxftype() == "MTEXT"
        d = _dialog(self.app, ("Новый " if nou else "Редактирование: ") + ("многострочный текст" if mtext else "текст"), 520, 380)
        d.columnconfigure(1, weight=1)
        tk.Label(d, text="Текст:", bg=CORP, fg=TEXT, font=F_MIC).grid(row=0, column=0, sticky="nw", padx=10, pady=6)
        tx = tk.Text(d, height=6, bg="#2b313a", fg=TEXT, insertbackground=TEXT, relief="flat", font=(F, -14), wrap="word")
        tx.grid(row=0, column=1, sticky="nsew", padx=10, pady=6); d.rowconfigure(0, weight=1)
        if not nou: tx.insert("1.0", e.plain_text() if mtext else e.dxf.text)
        h0 = (e.dxf.char_height if mtext else e.dxf.height) if not nou else self.inaltime
        r0 = (e.dxf.get("rotation", 0.0)) if not nou else 0.0
        eh = _rand(d, "Высота:", _intrare(round(h0, 4)), 1); er = _rand(d, "Поворот (°):", _intrare(round(r0, 3)), 2)
        straturi = sorted(l.dxf.name for l in self.doc.layers)
        cs = _rand(d, "Слой:", lambda p: ttk.Combobox(p, values=straturi, state="readonly"), 3); cs.set(e.dxf.layer if not nou else self.app.strat_curent())
        cc = _rand(d, "Цвет:", lambda p: ttk.Combobox(p, values=[n for n, _ in CULORI], state="readonly"), 4)
        cc.set(next((n for n, v in CULORI if v == (e.dxf.get("color", 256) if not nou else 256)), "ПоСлою"))
        tx.focus_set()
        def ok():
            t = tx.get("1.0", "end").rstrip("\n")
            if not t.strip(): return d.destroy()
            try: h = float(eh.get().replace(",", ".")); r = float(er.get().replace(",", "."))
            except ValueError: return messagebox.showerror("Текст", "Высота и поворот – числа.", parent=d)
            col = dict(CULORI)[cc.get()]
            self.app.ed.instantaneu(); msp = self.doc.modelspace()
            if nou:
                at = {"layer": cs.get(), "color": col}
                if mtext:
                    ent = msp.add_mtext(t.replace("\n", "\\P"), dxfattribs=dict(at, char_height=h, rotation=r)); ent.set_location(punct)
                else:
                    ent = msp.add_text(t.split("\n")[0], dxfattribs=dict(at, height=h, rotation=r)); ent.set_placement(punct)
                self.inaltime = h
            else:
                if mtext: e.text = t.replace("\n", "\\P"); e.dxf.char_height = h
                else: e.dxf.text = t.replace("\n", " "); e.dxf.height = h
                e.dxf.rotation = r; e.dxf.layer = cs.get(); e.dxf.color = col
            d.destroy(); self.app.ed.gata_modificare("Текст создан." if nou else "Текст изменен.")
        _butoane(d, ok, 5)

    # ---------------------------------------------------------------- editorul de atribute (dublu clic pe un bloc)
    def dialog_atribute(self, ins):
        atr = list(ins.attribs)
        if not atr: return self.dialog_proprietati([ins])
        d = _dialog(self.app, "Диспетчер атрибутов блоков – " + ins.dxf.name, 560, 120 + 34 * min(len(atr), 12)); d.columnconfigure(1, weight=1)
        tk.Label(d, text=f"Блок: {ins.dxf.name}    Слой: {ins.dxf.layer}", bg=CORP, fg=GRI, font=F_MIC).grid(row=0, column=0, columnspan=2, sticky="w", padx=10, pady=6)
        intr = []
        for i, a in enumerate(atr[:40]):
            intr.append((a, _rand(d, a.dxf.tag, _intrare(a.dxf.text), i + 1)))
        if intr: intr[0][1].focus_set()
        def ok():
            self.app.ed.instantaneu()
            for a, w in intr: a.dxf.text = w.get()
            d.destroy(); self.app.ed.gata_modificare("Атрибуты изменены.")
        _butoane(d, ok, len(intr) + 2)

    def dialog_proprietati(self, ents):
        if not ents: return self.app.mesaj("СВОЙСТВА: выберите объекты.")
        d = _dialog(self.app, "Свойства", 420, 200); d.columnconfigure(1, weight=1)
        tk.Label(d, text=f"Выбрано: {len(ents)}  ({', '.join(sorted({e.dxftype() for e in ents}))})", bg=CORP, fg=GRI, font=F_MIC).grid(row=0, column=0, columnspan=2, sticky="w", padx=10, pady=6)
        straturi = sorted(l.dxf.name for l in self.doc.layers)
        cs = _rand(d, "Слой:", lambda p: ttk.Combobox(p, values=straturi, state="readonly"), 1); cs.set(ents[0].dxf.layer)
        cc = _rand(d, "Цвет:", lambda p: ttk.Combobox(p, values=[n for n, _ in CULORI], state="readonly"), 2)
        cc.set(next((n for n, v in CULORI if v == ents[0].dxf.get("color", 256)), "ПоСлою"))
        def ok():
            self.app.ed.instantaneu(); col = dict(CULORI)[cc.get()]
            for e in ents: e.dxf.layer = cs.get(); e.dxf.color = col
            d.destroy(); self.app.ed.gata_modificare("Свойства изменены.")
        _butoane(d, ok, 3)

    # ---------------------------------------------------------------- toate textele din desen (TEXT, MTEXT, atribute)
    def texte(self):
        for e in self.doc.modelspace():
            t = e.dxftype()
            if t == "TEXT": yield e, e.dxf.text, e
            elif t == "MTEXT": yield e, e.text, e
            elif t == "INSERT":
                for a in e.attribs: yield a, a.dxf.text, e

    # ---------------------------------------------------------------- Найти и заменить
    def cauta_inlocuieste(self, initial=""):
        if not self.doc: return self.app.mesaj("Нет открытого чертежа.")
        d = tk.Toplevel(self.app.root); d.title("Найти и заменить"); d.configure(bg=CORP); d.transient(self.app.root); d.geometry("640x460")
        d.columnconfigure(1, weight=1)
        ec = _rand(d, "Найти:", _intrare(initial), 0); ei = _rand(d, "Заменить на:", _intrare(""), 1)
        vm = tk.BooleanVar(value=False); vc = tk.BooleanVar(value=False)
        o = tk.Frame(d, bg=CORP); o.grid(row=2, column=1, sticky="w", padx=10)
        for t, v in (("С учетом регистра", vm), ("Только слова целиком", vc)):
            tk.Checkbutton(o, text=t, variable=v, bg=CORP, fg=TEXT, selectcolor="#2b313a", activebackground=CORP, activeforeground=TEXT).pack(side="left", padx=(0, 12))
        tv = ttk.Treeview(d, columns=("t", "s", "b"), show="headings", style="Cad.Treeview", height=10)
        for c, t, w in (("t", "Текст", 330), ("s", "Слой", 120), ("b", "Объект", 130)): tv.heading(c, text=t); tv.column(c, width=w)
        tv.grid(row=3, column=0, columnspan=2, sticky="nsew", padx=10, pady=6); d.rowconfigure(3, weight=1)
        lbl = tk.Label(d, text="", bg=CORP, fg=GRI, font=F_MIC); lbl.grid(row=4, column=0, columnspan=2, sticky="w", padx=10)
        gasite = {}
        def model():
            s = ec.get()
            if not s: return None
            p = re.escape(s); p = r"\b" + p + r"\b" if vc.get() else p
            return re.compile(p, 0 if vm.get() else re.IGNORECASE)
        def cauta():
            tv.delete(*tv.get_children()); gasite.clear(); m = model()
            if not m: return
            for ent, txt, top in self.texte():
                if m.search(txt or ""):
                    iid = ent.dxf.handle; gasite[iid] = (ent, top)
                    tv.insert("", "end", iid=iid, values=(txt, top.dxf.layer, ("атрибут " + ent.dxf.tag + " в " + top.dxf.name) if ent.dxftype() == "ATTRIB" else ent.dxftype()))
            lbl.config(text=f"Найдено: {len(gasite)}")
            self.app.ed.sel = {t.dxf.handle for _, t in gasite.values()}; self.app.ed.arata_selectia()
        def inlocuieste(toate):
            m = model()
            if not m: return
            alese = list(gasite) if toate else list(tv.selection())
            if not alese: return lbl.config(text="Выберите строки в списке или «Заменить все».")
            self.app.ed.instantaneu(); n = 0
            for iid in alese:
                ent, top = gasite[iid]
                if ent.dxftype() == "MTEXT": ent.text = m.sub(ei.get(), ent.text)
                else: ent.dxf.text = m.sub(ei.get(), ent.dxf.text)
                n += 1
            self.app.ed.gata_modificare(f"Заменено: {n}"); cauta(); lbl.config(text=f"Заменено: {n}")
        def arata(e=None):
            s = tv.selection()
            if not s: return
            ent, top = gasite[s[0]]; b = self.app.v.bbox.get(top.dxf.handle)
            if b:
                v = self.app.v; W, H = v.wh(); v.cx, v.cy = (b[0] + b[2]) / 2, (b[1] + b[3]) / 2
                v.s = min(W / max(b[2] - b[0], 1e-6), H / max(b[3] - b[1], 1e-6)) * 0.15; v.deseneaza()
            self.app.ed.sel = {top.dxf.handle}; self.app.ed.arata_selectia()
        tv.bind("<Double-1>", arata)
        f = tk.Frame(d, bg=CORP); f.grid(row=5, column=0, columnspan=2, sticky="e", padx=10, pady=8)
        for t, cmd in (("Найти", cauta), ("Заменить", lambda: inlocuieste(False)), ("Заменить все", lambda: inlocuieste(True)), ("Показать", arata), ("Готово", d.destroy)):
            b = tk.Label(f, text=t, bg=ALBASTRU if t == "Найти" else FILE, fg="white", font=F_MIC, padx=12, pady=4); b.pack(side="left", padx=3)
            b.bind("<Button-1>", lambda e, c=cmd: c())
        ec.bind("<Return>", lambda e: cauta()); ec.focus_set()
        if initial: cauta()

    # ---------------------------------------------------------------- numerotare automată (ca Auto Number / TCOUNT)
    def numerotare(self):
        ed = self.app.ed
        ents = [ed.ent(h) for h in ed.sel]; ents = [e for e in ents if e is not None and e.dxftype() in ("TEXT", "MTEXT", "INSERT")]
        if not ents: return self.app.mesaj("Auto Number: выберите тексты или блоки с атрибутами (рамкой), затем команду.")
        etichete = sorted({a.dxf.tag for e in ents if e.dxftype() == "INSERT" for a in e.attribs})
        d = _dialog(self.app, "Auto Number – автонумерация", 470, 330); d.columnconfigure(1, weight=1)
        tk.Label(d, text=f"Выбрано: {len(ents)} (тексты и блоки)", bg=CORP, fg=GRI, font=F_MIC).grid(row=0, column=0, columnspan=2, sticky="w", padx=10, pady=6)
        cs = _rand(d, "Порядок:", lambda p: ttk.Combobox(p, values=["по X (слева направо)", "по Y (сверху вниз)", "по Y (снизу вверх)", "по X (справа налево)"], state="readonly"), 1); cs.current(0)
        e0 = _rand(d, "Начальный номер:", _intrare(1), 2); ep = _rand(d, "Шаг:", _intrare(1), 3)
        cm = _rand(d, "Способ:", lambda p: ttk.Combobox(p, values=["Заменить текст", "Префикс (номер + текст)", "Суффикс (текст + номер)"], state="readonly"), 4); cm.current(0)
        ca = _rand(d, "Атрибут блока:", lambda p: ttk.Combobox(p, values=etichete, state="readonly"), 5)
        if etichete: ca.set(next((t for t in etichete if t.upper() in ("NR", "NUM", "НОМЕР", "N", "NUMAR")), etichete[0]))
        ef = _rand(d, "Текст до / после:", _intrare(""), 6)
        def poz(e):
            p = e.dxf.insert; return p.x, p.y
        def ok():
            try: n = float(e0.get().replace(",", ".")); pas = float(ep.get().replace(",", "."))
            except ValueError: return messagebox.showerror("Auto Number", "Номер и шаг – числа.", parent=d)
            k = cs.current(); cheie = {0: lambda e: (poz(e)[0], -poz(e)[1]), 1: lambda e: (-poz(e)[1], poz(e)[0]), 2: lambda e: (poz(e)[1], poz(e)[0]), 3: lambda e: (-poz(e)[0], -poz(e)[1])}[k]
            ed.instantaneu(); txt_f = ef.get()
            for e in sorted(ents, key=cheie):
                nr = str(int(n)) if float(n).is_integer() else str(n)
                if e.dxftype() == "INSERT":
                    a = next((a for a in e.attribs if a.dxf.tag == ca.get()), None)
                    if a is None: continue
                    tinta, vechi = a, a.dxf.text
                else: tinta, vechi = e, (e.text if e.dxftype() == "MTEXT" else e.dxf.text)
                nou = {0: txt_f + nr if txt_f else nr, 1: nr + txt_f + vechi, 2: vechi + txt_f + nr}[cm.current()]
                if tinta.dxftype() == "MTEXT": tinta.text = nou
                else: tinta.dxf.text = nou
                n += pas
            d.destroy(); ed.gata_modificare(f"Пронумеровано: {len(ents)}")
        _butoane(d, ok, 7)

    # ---------------------------------------------------------------- atribute ↔ Excel
    def export_excel(self):
        if not self.doc: return self.app.mesaj("Нет открытого чертежа.")
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        sel = self.app.ed.sel
        blocuri = [e for e in self.doc.modelspace() if e.dxftype() == "INSERT" and e.attribs and (not sel or e.dxf.handle in sel)]
        texte = [e for e in self.doc.modelspace() if e.dxftype() in ("TEXT", "MTEXT") and (not sel or e.dxf.handle in sel)]
        if not blocuri and not texte: return self.app.mesaj("Нет блоков с атрибутами или текстов.")
        baza = os.path.splitext(os.path.basename(self.app.cale or "Чертеж1"))[0]
        p = filedialog.asksaveasfilename(title="Экспорт атрибутов в Excel", initialfile=baza + "_atribute.xlsx", defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if not p: return
        wb = Workbook(); B = Font(bold=True, color="FFFFFF"); Fl = PatternFill("solid", fgColor="2B5797")
        ws = wb.active; ws.title = "Атрибуты"
        tags = []
        for e in blocuri:
            for a in e.attribs:
                if a.dxf.tag not in tags: tags.append(a.dxf.tag)
        ws.append(["HANDLE", "БЛОК", "СЛОЙ", "X", "Y"] + tags)
        for e in blocuri:
            v = {a.dxf.tag: a.dxf.text for a in e.attribs}
            ws.append(["'" + e.dxf.handle, e.dxf.name, e.dxf.layer, round(e.dxf.insert.x, 3), round(e.dxf.insert.y, 3)] + [v.get(t, "") for t in tags])
        ws2 = wb.create_sheet("Тексты"); ws2.append(["HANDLE", "ТИП", "СЛОЙ", "X", "Y", "ТЕКСТ"])
        for e in texte:
            ws2.append(["'" + e.dxf.handle, e.dxftype(), e.dxf.layer, round(e.dxf.insert.x, 3), round(e.dxf.insert.y, 3), e.plain_text() if e.dxftype() == "MTEXT" else e.dxf.text])
        for w in (ws, ws2):
            for c in w[1]: c.font = B; c.fill = Fl
            w.freeze_panes = "A2"
            for col in w.columns: w.column_dimensions[col[0].column_letter].width = max(10, min(40, max(len(str(c.value or "")) for c in col) + 2))
        wb.save(p); self.app.mesaj(f"Экспорт: {len(blocuri)} блоков, {len(texte)} текстов → {p}")
        if os.name == "nt":
            try: os.startfile(p)
            except Exception: pass

    def import_excel(self):
        if not self.doc: return self.app.mesaj("Нет открытого чертежа.")
        from openpyxl import load_workbook
        p = filedialog.askopenfilename(title="Импорт атрибутов из Excel", filetypes=[("Excel", "*.xlsx")])
        if not p: return
        wb = load_workbook(p, data_only=True); self.app.ed.instantaneu(); n = 0; t = 0
        if "Атрибуты" in wb.sheetnames:
            ws = wb["Атрибуты"]; cap = [str(c.value or "") for c in ws[1]]
            for r in ws.iter_rows(min_row=2, values_only=True):
                h = str(r[0] or "").lstrip("'"); e = self.app.ed.ent(h)
                if e is None or e.dxftype() != "INSERT": continue
                for a in e.attribs:
                    if a.dxf.tag in cap:
                        val = r[cap.index(a.dxf.tag)]; val = "" if val is None else (str(int(val)) if isinstance(val, float) and val.is_integer() else str(val))
                        if a.dxf.text != val: a.dxf.text = val; n += 1
        if "Тексты" in wb.sheetnames:
            for r in wb["Тексты"].iter_rows(min_row=2, values_only=True):
                h = str(r[0] or "").lstrip("'"); e = self.app.ed.ent(h)
                if e is None or r[5] is None: continue
                val = str(r[5])
                if e.dxftype() == "TEXT" and e.dxf.text != val: e.dxf.text = val; t += 1
                elif e.dxftype() == "MTEXT" and e.plain_text() != val: e.text = val.replace("\n", "\\P"); t += 1
        self.app.ed.gata_modificare(f"Импорт из Excel: изменено атрибутов {n}, текстов {t}.")
