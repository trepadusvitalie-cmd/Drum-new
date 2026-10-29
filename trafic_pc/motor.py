"""Motorul analizei: urmărirea obiectelor între cadre, porțile de numărare, zonele, viteza și exportul Excel.
Rulare fără fereastră:  python motor.py film.mp4 config.json  ->  film_trafic.xlsx (+ film_trafic.mp4 cu desenele)"""
import json, math, os, sys, time, datetime
from collections import Counter, defaultdict
import numpy as np
import cv2
import detector, raport
from detector import CLASE, NUME, CUL, log


# ---------------------------------------------------------------- urmărirea
def iou(a, b):
    l, t, r, bo = max(a[0], b[0]), max(a[1], b[1]), min(a[2], b[2]), min(a[3], b[3])
    i = max(0.0, r - l) * max(0.0, bo - t)
    return i / ((a[2] - a[0]) * (a[3] - a[1]) + (b[2] - b[0]) * (b[3] - b[1]) - i + 1e-9)


class Urma:
    _n = 0

    ANCORA = "centru"                                        # "centru" (camera de sus) sau "jos" (camera din lateral)

    def __init__(self, box, scor, cls, t):
        Urma._n += 1; self.id = Urma._n
        self.box = np.array(box, float); self.v = np.zeros(4); self.t = t; self.tn = t
        self.voturi = Counter({cls: scor}); self.loviri = 1; self.pierdut = 0
        self.drum = [(t, *self.punct())]                     # traiectoria: (timp, x, y) – punctul de contact cu solul
        self.p0 = self.drum[0][1:]                           # de unde a pornit (drumul se scurtează la 300 de puncte)
        self.porti = set(); self.zone = {}; self.viteza = None; self.verif = 0     # până unde s-a verificat drumul la porți
        self.marimi = []; self.loviri_ai = 1; self._stat(box)

    @property
    def cls(self): return self.voturi.most_common(1)[0][0]
    @property
    def sigur(self): return self.loviri >= 3

    def punct(self):
        b = self.box; return ((b[0] + b[2]) / 2, b[3] if Urma.ANCORA == "jos" else (b[1] + b[3]) / 2)

    def _stat(self, box):
        w, h = box[2] - box[0], box[3] - box[1]
        if w > 2 and h > 2: self.marimi.append(((box[1] + box[3]) / 2, w * h, h / w)); self.marimi = self.marimi[-60:]

    def prezis(self, t):
        return self.box + self.v * (t - self.t)

    def actualizeaza(self, box, scor, cls, t):
        box = np.array(box, float); dt = max(1e-3, t - self.t)
        self.v = 0.6 * self.v + 0.4 * (box - self.box) / dt
        self.box = box; self.t = t; self.voturi[cls] += scor; self.loviri += 1; self.pierdut = 0
        if scor > 0: self.loviri_ai += 1; self._stat(box)
        self.drum.append((t, *self.punct()))
        if len(self.drum) > 300: self.verif -= len(self.drum) - 300; self.drum = self.drum[-300:]


class Urmaritor:
    """Potrivire după suprapunere (IoU) cu poziția prezisă – întâi detecțiile sigure, apoi cele slabe (ca ByteTrack)."""

    def __init__(self, max_pierdut_s=1.5):
        self.urme = []; self.max_pierdut_s = max_pierdut_s

    def pas(self, dets, t, pete=()):
        tari = [d for d in dets if d[4] >= 0.45]; slabe = [d for d in dets if d[4] < 0.45]
        libere = list(self.urme)
        for grup, prag in ((tari, 0.25), (slabe, 0.35)):
            per = []
            for k, d in enumerate(grup):
                for u in libere:
                    s = iou(u.prezis(t), d[:4])
                    if s >= prag: per.append((s, k, u))
            per.sort(key=lambda x: -x[0]); fol_d, fol_u = set(), set()
            for s, k, u in per:
                if k in fol_d or u.id in fol_u: continue
                fol_d.add(k); fol_u.add(u.id); d = grup[k]; u.actualizeaza(d[:4], d[4], d[5], t)
            libere = [u for u in libere if u.id not in fol_u]
            # a doua șansă pentru obiectele rapide: centrul prezis aproape de centrul detecției
            per = []
            for k, d in enumerate(grup):
                if k in fol_d: continue
                cd = np.array([(d[0] + d[2]) / 2, (d[1] + d[3]) / 2]); sz = max(d[2] - d[0], d[3] - d[1])
                for u in libere:
                    pb = u.prezis(t); cp = np.array([(pb[0] + pb[2]) / 2, (pb[1] + pb[3]) / 2])
                    dist = float(np.hypot(*(cd - cp)))
                    if dist < 0.7 * max(sz, pb[2] - pb[0], pb[3] - pb[1]) and (u.cls == d[5] or d[4] < 0.6): per.append((dist, k, u))
            per.sort(key=lambda x: x[0])
            for dist, k, u in per:
                if k in fol_d or u.id in fol_u: continue
                fol_d.add(k); fol_u.add(u.id); d = grup[k]; u.actualizeaza(d[:4], d[4], d[5], t)
            libere = [u for u in libere if u.id not in fol_u]
            if grup is tari:
                for k, d in enumerate(grup):
                    if k not in fol_d: self.urme.append(Urma(d[:4], d[4], d[5], t))
        # urmele încă nepotrivite continuă pe pata de mișcare de lângă poziția prezisă (AI-ul a pierdut obiectul)
        luate = set()
        for u in sorted(libere, key=lambda u: -u.loviri):
            if not u.sigur: continue
            pb = u.prezis(t); best = None
            for k, b in enumerate(pete):
                if k in luate: continue
                s_ = iou(pb, b)
                if s_ > 0.25 and (best is None or s_ > best[0]): best = (s_, k)
            if best:
                b = np.array(pete[best[1]], float); luate.add(best[1])
                # păstrăm mărimea obiectului, mutăm doar centrul (pata poate fi mai mare/mică)
                c = (b[:2] + b[2:]) / 2; wh = (u.box[2:] - u.box[:2]) * 0.5 + (b[2:] - b[:2]) * 0.5
                u.actualizeaza(np.r_[c - wh / 2, c + wh / 2], 0.0, u.cls, t); u.loviri -= 1; u.pe_miscare = True
        for u in libere:
            if u.t != t: u.pierdut += 1
        self.urme = [u for u in self.urme if t - u.t <= self.max_pierdut_s]
        return [u for u in self.urme if u.sigur and u.t == t]


class Miscare:
    """Petele care se mișcă (scăderea fundalului) – rezervă pentru cadrele în care AI-ul nu vede obiectul."""

    def __init__(self):
        self.bg = cv2.createBackgroundSubtractorMOG2(history=300, varThreshold=32, detectShadows=True); self.k = None

    def __call__(self, img):
        h, w = img.shape[:2]; f = 480 / max(h, w)
        sm = cv2.resize(img, (int(w * f), int(h * f)), interpolation=cv2.INTER_AREA)
        m = self.bg.apply(sm); m = (m == 255).astype(np.uint8) * 255                 # fără umbre (127)
        m = cv2.morphologyEx(m, cv2.MORPH_OPEN, np.ones((3, 3), np.uint8))
        m = cv2.dilate(m, np.ones((7, 7), np.uint8), iterations=2)
        n, lab, st, _ = cv2.connectedComponentsWithStats(m)
        out = []
        for i in range(1, n):
            x, y, ww, hh, a = st[i]
            if a < 0.0015 * m.size: continue
            out.append((x / f, y / f, (x + ww) / f, (y + hh) / f))
        return out


# ---------------------------------------------------------------- geometrie
def taie(p, q, a, b):
    """segmentul p->q taie poarta a->b? -> +1 / -1 după sens, 0 dacă nu"""
    def cr(o, x, y): return (x[0] - o[0]) * (y[1] - o[1]) - (x[1] - o[1]) * (y[0] - o[0])
    d1, d2 = cr(a, b, p), cr(a, b, q); d3, d4 = cr(p, q, a), cr(p, q, b)
    if (d1 > 0) != (d2 > 0) and (d3 > 0) != (d4 > 0) and d1 != d2:
        return 1 if d1 > 0 else -1
    return 0


def in_poligon(pt, poly):
    return cv2.pointPolygonTest(np.array(poly, np.float32), (float(pt[0]), float(pt[1])), False) >= 0


class Calibrare:
    """4 puncte în imagine = colțurile unui dreptunghi real (lățime × lungime în metri) -> metri pe sol."""

    def __init__(self, c):
        src = np.array(c["puncte"], np.float32); L, W = float(c["lungime_m"]), float(c["latime_m"])
        dst = np.array([[0, 0], [W, 0], [W, L], [0, L]], np.float32)
        self.H = cv2.getPerspectiveTransform(src, dst)

    def metri(self, x, y):
        p = self.H @ np.array([x, y, 1.0]); return p[:2] / p[2]

    def viteza(self, drum, fereastra=1.0):
        """km/h din ultima secundă a traiectoriei"""
        if len(drum) < 3: return None
        t1 = drum[-1][0]; pts = [d for d in drum if d[0] >= t1 - fereastra]
        if len(pts) < 3 or pts[-1][0] - pts[0][0] < 0.3: return None
        a, b = self.metri(*pts[0][1:]), self.metri(*pts[-1][1:])
        return float(np.hypot(*(b - a)) / (pts[-1][0] - pts[0][0]) * 3.6)


# ---------------------------------------------------------------- analiza
TOATE = "Toate vehiculele"                  # „poarta” folosită când nu e desenată nicio poartă


class Analiza:
    def __init__(self, cfg):
        self.cfg = cfg
        self.porti = cfg.get("porti", []); self.zone = cfg.get("zone", [])
        self.cal = Calibrare(cfg["calibrare"]) if cfg.get("calibrare") else None
        Urma.ANCORA = cfg.get("ancora", "centru")
        self.urm = Urmaritor(); self.ref = raport.Referinta()
        self.clase = set(cfg.get("clase", ["bicicleta", "masina", "motocicleta", "autobuz", "tramvai", "camion"]))
        self.treceri = []                       # (t, poarta, sens, clasa, id, viteza)
        self.zone_ev = []                       # (id, clasa, zona, t_intrare, t_iesire, alerta)
        self.alerte = []                        # (t, zona, id, clasa, motiv)
        self.contor = {p["nume"]: {s: Counter() for s in ("A→B", "B→A")} for p in self.porti}
        # fără porți: fiecare vehicul care s-a deplasat se numără o dată, când iese din imagine (sens: în jos / în sus)
        self.fara_porti = not self.porti
        if self.fara_porti: self.contor = {TOATE: {s: Counter() for s in ("A→B", "B→A")}}
        self.vazute = {}
        self.in_zona = {z["nume"]: {} for z in self.zone}

    def pas(self, dets, t, pete=()):
        dets = [d for d in dets if d[5] in self.clase]
        for d in dets:                                           # mărimea tipică a unei mașini, pe înălțimea imaginii
            if d[5] == "masina" and d[4] > 0.6: self.ref.adauga((d[1] + d[3]) / 2, (d[2] - d[0]) * (d[3] - d[1]))
        active = self.urm.pas(dets, t, pete)
        if self.fara_porti:
            for u in active: self.vazute[u.id] = u
            vii = {u.id for u in self.urm.urme}
            for uid in [i for i in self.vazute if i not in vii]: self._numara(self.vazute.pop(uid))
        for u in active:
            if self.cal: u.viteza = self.cal.viteza(u.drum)
            # toate segmentele drumului încă neverificate (și cele de dinainte ca urma să devină sigură)
            for k in range(max(1, u.verif), len(u.drum)):
                p, q = u.drum[k - 1][1:], u.drum[k][1:]
                for g in self.porti:
                    s = taie(p, q, g["a"], g["b"])
                    if s and g["nume"] not in u.porti:
                        u.porti.add(g["nume"]); sens = "A→B" if s > 0 else "B→A"
                        self.contor[g["nume"]][sens][u.cls] += 1
                        self.treceri.append((u.drum[k][0], g["nume"], sens, u.cls, u.id, u.viteza, raport.clasa_raport(u, self.ref)))
            u.verif = len(u.drum)
            for z in self.zone:
                inside = in_poligon(u.drum[-1][1:], z["poligon"]); zn = z["nume"]; dentro = self.in_zona[zn]
                if inside and u.id not in dentro: dentro[u.id] = [t, u.cls, False]
                if inside:
                    st = float(z.get("stationare_s", 0) or 0)
                    clase_alerta = z.get("alerta_clase", [])
                    if not dentro[u.id][2] and ((st and t - dentro[u.id][0] >= st) or u.cls in clase_alerta):
                        dentro[u.id][2] = True
                        motiv = f"staționează peste {st:.0f} s" if st and t - dentro[u.id][0] >= st else "prezență în zonă"
                        self.alerte.append((t, zn, u.id, u.cls, motiv))
        # ieșirile din zone (inclusiv urmele pierdute)
        vii = {u.id: u for u in self.urm.urme}
        for z in self.zone:
            dentro = self.in_zona[z["nume"]]
            for uid in list(dentro):
                u = vii.get(uid)
                if u is None or (u.t == t and not in_poligon(u.drum[-1][1:], z["poligon"])):
                    t0, cls, al = dentro.pop(uid); t1 = u.t if u else t
                    self.zone_ev.append((uid, cls, z["nume"], t0, t1, al))
        return active

    def _numara(self, u):
        """numărarea fără porți: doar obiectele care s-au deplasat cel puțin o înălțime de chenar (nu cele parcate)"""
        x0, y0 = u.p0; x1, y1 = u.drum[-1][1:]; h = max(1.0, u.box[3] - u.box[1])
        if math.hypot(x1 - x0, y1 - y0) < h: return
        sens = "A→B" if y1 > y0 else "B→A"
        self.contor[TOATE][sens][u.cls] += 1
        self.treceri.append((u.t, TOATE, sens, u.cls, u.id, u.viteza, raport.clasa_raport(u, self.ref)))

    def final(self, t):
        if self.fara_porti:
            for u in list(self.vazute.values()): self._numara(u)
            self.vazute = {}; self.treceri.sort(key=lambda x: x[0])
        for z in self.zone:
            for uid, (t0, cls, al) in self.in_zona[z["nume"]].items(): self.zone_ev.append((uid, cls, z["nume"], t0, t, al))
            self.in_zona[z["nume"]] = {}


# ---------------------------------------------------------------- desenul pe cadru
def deseneaza(img, an, active, t=None, ora=None):
    v = img.copy(); th = max(2, img.shape[1] // 500); fs = max(0.5, img.shape[1] / 1600)
    for z in an.zone:
        p = np.array(z["poligon"], np.int32); ov = v.copy(); cv2.fillPoly(ov, [p], (0, 200, 255)); v = cv2.addWeighted(ov, 0.15, v, 0.85, 0)
        cv2.polylines(v, [p], True, (0, 170, 255), th); cv2.putText(v, z["nume"], tuple(p[0]), 0, fs, (0, 120, 220), th)
    for g in an.porti:
        a, b = tuple(map(int, g["a"])), tuple(map(int, g["b"])); cv2.line(v, a, b, (60, 220, 60), th + 1)
        c = an.contor[g["nume"]]; n1, n2 = sum(c["A→B"].values()), sum(c["B→A"].values())
        cv2.putText(v, f"{g['nume']}  A>B {n1}  B>A {n2}", (a[0], a[1] - 8), 0, fs, (30, 150, 30), th)
    for u in active:
        x0, y0, x1, y1 = map(int, u.box); c = CUL[u.cls]
        cv2.rectangle(v, (x0, y0), (x1, y1), c, th)
        pts = np.array([d[1:] for d in u.drum[-60:]], np.int32); cv2.polylines(v, [pts], False, c, max(1, th - 1))
        t_ = f"{u.id} {NUME[u.cls]}" + (f" {u.viteza:.0f} km/h" if u.viteza else "")
        (tw, tht), _ = cv2.getTextSize(t_, 0, fs * 0.8, max(1, th - 1))
        cv2.rectangle(v, (x0, y0 - tht - 8), (x0 + tw + 6, y0), c, -1)
        cv2.putText(v, t_, (x0 + 3, y0 - 4), 0, fs * 0.8, (255, 255, 255), max(1, th - 1), cv2.LINE_AA)
    if ora: cv2.putText(v, ora, (10, int(30 * fs * 1.6)), 0, fs * 1.1, (255, 255, 255), th + 1, cv2.LINE_AA)
    return v


# ---------------------------------------------------------------- exportul Excel
def ora_txt(t, start):
    return (start + datetime.timedelta(seconds=t)).strftime("%Y-%m-%d %H:%M:%S") if start else f"{int(t // 60):02d}:{t % 60:04.1f}"


def export_excel(an, path, start=None, durata=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); B = Font(bold=True, color="FFFFFF"); F = PatternFill("solid", fgColor="17324D")
    def antet(ws, cols):
        ws.append(cols)
        for c in ws[1]: c.font = B; c.fill = F; c.alignment = Alignment(horizontal="center")
        for k, col in enumerate(cols): ws.column_dimensions[chr(65 + k)].width = max(12, len(col) + 4)
        ws.freeze_panes = "A2"
    # rezumat pe porți
    ws = wb.active; ws.title = "Rezumat porți"
    antet(ws, ["Poarta", "Sens"] + raport.RAPORT + ["Total"])
    for g in an.contor:
        for s in ("A→B", "B→A"):
            c = Counter(x[6] for x in an.treceri if x[1] == g and x[2] == s and x[6])
            ws.append([g, s] + [c.get(k, 0) for k in raport.RAPORT] + [sum(c.values())])
    # fiecare trecere
    ws = wb.create_sheet("Treceri")
    antet(ws, ["Ora", "Timp video (s)", "Poarta", "Sens", "Clasa raport", "Clasa AI", "ID", "Viteza km/h"])
    for t, g, s, c, i, v, rc in an.treceri: ws.append([ora_txt(t, start), round(t, 1), g, s, rc or "", NUME[c], i, round(v, 1) if v else None])
    # pe intervale
    for minute in (5, 15, 60):
        ws = wb.create_sheet(f"Intervale {minute} min")
        antet(ws, ["Început interval", "Poarta", "Sens"] + raport.RAPORT + ["Total", "Viteza medie km/h"])
        L = minute * 60; tmax = max([durata or 0] + [x[0] for x in an.treceri])
        # intervale la ore rotunde (08:00, 08:15…): decalajul începutului filmării față de ora rotundă
        d0 = ((start.minute * 60 + start.second) % L) if start else 0
        n = int((tmax + d0) // L) + 1
        for k in range(n):
            a_, b_ = k * L - d0, (k + 1) * L - d0
            for g in an.contor:
                for s in ("A→B", "B→A"):
                    sel = [x for x in an.treceri if x[1] == g and x[2] == s and a_ <= x[0] < b_]
                    sel = [x for x in sel if x[6]]; c = Counter(x[6] for x in sel); vs = [x[5] for x in sel if x[5]]
                    ws.append([ora_txt(a_, start), g, s] + [c.get(q, 0) for q in raport.RAPORT] + [len(sel), round(float(np.mean(vs)), 1) if vs else None])
    # zone
    ws = wb.create_sheet("Zone")
    antet(ws, ["Zona", "ID", "Clasa", "Intrare", "Ieșire", "Durata (s)", "Alertă"])
    for uid, c, z, t0, t1, al in sorted(an.zone_ev, key=lambda x: x[3]):
        ws.append([z, uid, NUME[c], ora_txt(t0, start), ora_txt(t1, start), round(t1 - t0, 1), "DA" if al else ""])
    ws = wb.create_sheet("Alerte")
    antet(ws, ["Ora", "Zona", "ID", "Clasa", "Motiv"])
    for t, z, i, c, m in an.alerte: ws.append([ora_txt(t, start), z, i, NUME[c], m])
    if start:
        rand = [(start + datetime.timedelta(seconds=x[0]), x[1], x[2], x[6], x[3], x[4], x[5]) for x in an.treceri]
        raport.foaie_zilnica(wb, rand, luni_intregi=False)
        raport.scrie_csv(os.path.splitext(path)[0].replace("_trafic", "") + "_treceri.csv", rand)
    wb.save(path)


# ---------------------------------------------------------------- rularea pe un film
def ruleaza(video, cfg, out_xlsx=None, out_mp4=None, det=None, progres=None, cadru_cb=None, stop=None, pauza=None):
    """Analizează filmul (sau camera rtsp://…). progres(fracție, text); cadru_cb(imagine desenată, analiza)."""
    det = det or detector.Detector(cfg.get("model", "yolox_s.onnx"), cfg.get("prag", 0.3))
    cap = cv2.VideoCapture(video)
    if not cap.isOpened(): raise RuntimeError(f"nu se poate deschide {video}")
    fps = cap.get(cv2.CAP_PROP_FPS) or 25.0; nfr = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
    live = str(video).lower().startswith(("rtsp:", "http:", "https:")) or nfr <= 0
    pas = max(1, int(round(fps / float(cfg.get("fps_analiza", 10)))))
    start = None
    if cfg.get("ora_start"):
        try: start = datetime.datetime.strptime(cfg["ora_start"], "%Y-%m-%d %H:%M:%S")
        except ValueError: start = None
    if live and start is None: start = datetime.datetime.now()
    an = Analiza(cfg); vw = None; fi = -1; t = 0.0; t0 = time.time(); nproc = 0; t_cb = 0.0
    misc = Miscare() if cfg.get("miscare", True) else None
    while True:
        if stop and stop(): break
        while pauza and pauza() and not (stop and stop()): time.sleep(0.1)
        ok = cap.grab(); fi += 1
        if not ok: break
        if fi % pas: continue
        ok, img = cap.retrieve()
        if not ok: break
        t = (time.time() - t0) if live else fi / fps
        dets = det(img); active = an.pas(dets, t, misc(img) if misc else ()); nproc += 1
        vrea_cb = cadru_cb and time.time() - t_cb >= 0.2
        if out_mp4 or vrea_cb:
            v = deseneaza(img, an, active, t, ora_txt(t, start))
            if out_mp4:
                if vw is None: vw = cv2.VideoWriter(out_mp4, cv2.VideoWriter_fourcc(*"mp4v"), fps / pas, (v.shape[1], v.shape[0]))
                vw.write(v)
            if vrea_cb: cadru_cb(v, an); t_cb = time.time()
        if progres and nproc % 5 == 0:
            progres(fi / nfr if nfr > 0 else 0.0, f"{ora_txt(t, start)} · {nproc / max(1e-3, time.time() - t0):.1f} cadre/s")
    an.final(t); cap.release()
    if vw: vw.release()
    if out_xlsx: export_excel(an, out_xlsx, start, t)
    return an


def main():
    if len(sys.argv) < 3:
        print(__doc__); return
    video, cfgp = sys.argv[1], sys.argv[2]; cfg = json.load(open(cfgp, encoding="utf-8"))
    baza = os.path.splitext(video)[0] + "_trafic"
    an = ruleaza(video, cfg, baza + ".xlsx", baza + ".mp4", progres=lambda f, s: log(f"{f * 100:5.1f}%  {s}"))
    for g, ss in an.contor.items():
        for s, c in ss.items(): log(f"{g} {s}: " + ", ".join(f"{NUME[k]} {v}" for k, v in c.items()) or "0")
    log(f"Zone: {len(an.zone_ev)} treceri prin zone, {len(an.alerte)} alerte -> {baza}.xlsx")


if __name__ == "__main__":
    main()
